"""Source-preserving B/C/F tyre statements and month-end engine/product matrices.

These tabs have a separate durable cache. Existing KPI and vehicle-history
calculations are not inputs to these source snapshots.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
from calendar import monthrange
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup

CACHE_TITLE = "_REPORT_DETAILS_HISTORY"
SCHEMA = "report-details-1"
TYRE_TITLE = "Tyre Statements"
ENGINE_TITLE = "Engine & Product KMPL"
TYRE_INPUT_TITLE = "_TYRE_INPUTS"
TAB_TITLES = (TYRE_TITLE, ENGINE_TITLE)
FIRST_PERIOD = "2025-04"
BASE = "http://103.44.14.20"
IST = ZoneInfo("Asia/Kolkata")
NAVY, BLUE, INK, PALE, LINE = "12345B", "1976B9", "23364D", "F3F6FA", "DCE5EF"
STATEMENTS = {
    "B": "Mechanical defects & premature tyre failures — Statement B",
    "C": "Scrap tyres — Statement C",
    "F": "Tyres below 2 mm, stones & worn smooth — Statement F",
}
SCRAP_GROUPS = ("New Scrap Tyres", "Ist RC Scrap Tyres", "2nd RC Scrap Tyres", "3rd RC and above Scrap Tyres")


def scrap_table(html):
    """C has two header levels: retain stage identity for repeated S1..S9 names."""
    soup, choices = BeautifulSoup(html, "html.parser"), []
    for table in soup.find_all("table"):
        trs = table.find_all("tr")
        if len(trs) < 3:
            continue
        first = trs[0].find_all("th", recursive=False)
        second = trs[1].find_all("th", recursive=False)
        if [key(c.get_text(" ", strip=True)) for c in first[7:]] != list(map(key, SCRAP_GROUPS)):
            continue
        expected = [f"S{i}" for i in range(1, 10)] + ["TOTAL"]
        if (len(first) != 11 or any(c.get("rowspan") != "2" for c in first[:7])
                or any(c.get("colspan") != "10" for c in first[7:])
                or [key(c.get_text(" ", strip=True)) for c in second] != expected * 4):
            raise ValueError("Statement C grouped headings are not recognised")
        headers = [c.get_text(" ", strip=True) for c in first[:7]]
        if list(map(key, headers)) != ["SNUM", "MONTHYEAR", "ZONENAME", "REGION", "DEPOT", "TOTALSCRAP", "TYRESIZE"]:
            raise ValueError("Statement C identifying headings are not recognised")
        headers += [f"{stage} — {c.get_text(' ', strip=True)}"
                    for i, stage in enumerate(SCRAP_GROUPS) for c in second[i*10:(i+1)*10]]
        rows = table_rows(table)[2:]
        if any(len(row) != 47 for row in rows):
            raise ValueError("Statement C row width does not match grouped headings")
        choices.append((headers, rows))
    if len(choices) != 1:
        raise ValueError("Expected one recognised Statement C table")
    return choices[0]


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip().upper()


def key(value):
    return re.sub(r"[^A-Z0-9]", "", norm(value))


def month_end(period):
    day = datetime.strptime(period, "%Y-%m").date()
    if day.strftime("%Y-%m") != period:
        raise ValueError("Report month must be YYYY-MM")
    return day.replace(day=monthrange(day.year, day.month)[1])


def fy_start(period):
    day = month_end(period)
    return date(day.year if day.month >= 4 else day.year - 1, 4, 1)


def fy_label(period):
    year = fy_start(period).year
    return f"{year}-{str(year + 1)[-2:]}"


def annual_periods(selected):
    """All source months from FY25-26 through the chosen month, never future views."""
    endpoint = month_end(selected)
    out, year, month = [], 2025, 4
    while date(year, month, 1) <= endpoint:
        out.append(f"{year:04d}-{month:02d}")
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return out


def effective_end(period, today=None):
    """Open months use completed days and are explicitly provisional."""
    today = today or datetime.now(IST).date()
    last = min(month_end(period), today - timedelta(days=1))
    if last.strftime("%Y-%m") != period:
        raise ValueError("No completed source day exists in the requested month")
    return last


def number(text):
    value = str("" if text is None else text).strip().replace(",", "")
    if value in ("", "-", "—", "–"):
        return None
    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?%?", value):
        raise ValueError(f"Non-numeric source metric: {text!r}")
    result = float(value.rstrip("%"))
    if not math.isfinite(result) or result < 0:
        raise ValueError("Source metric must be finite and non-negative")
    return int(result) if result.is_integer() else result


def table_rows(table):
    """Expand source colspans: TOTAL's two-cell label must not shift product totals."""
    rows = []
    for row in table.find_all("tr"):
        values = []
        for cell in row.find_all(["th", "td"], recursive=False):
            text = cell.get_text(" ", strip=True)
            span = int(cell.get("colspan", 1))
            if span < 1 or span > 100:
                raise ValueError("Unsupported source colspan")
            values.extend([text] + [""] * (span - 1))
        if values:
            rows.append(values)
    return rows


def source_table(html, required):
    soup = BeautifulSoup(html, "html.parser")
    choices = []
    for table in soup.find_all("table"):
        rows = table_rows(table)
        if rows and all(key(label) in {key(x) for x in rows[0]} for label in required):
            choices.append(rows)
    if len(choices) != 1:
        raise ValueError("Expected one recognised source detail table; login/empty pages are unavailable")
    rows = choices[0]
    headers = rows[0]
    if len(set(map(key, headers))) != len(headers) or any(not key(x) for x in headers):
        raise ValueError("Ambiguous source column headings")
    if any(len(row) != len(headers) for row in rows[1:]):
        raise ValueError("Source row width does not match the headings")
    return soup, headers, rows[1:]


def parse_tyre(html, statement, depot_code, period, depot):
    if statement not in STATEMENTS:
        raise ValueError("Unknown tyre statement")
    required = ["Month-Year", "Depot", "Tyre size",
                "Mismatching" if statement == "B" else "Worn Smooth %"]
    if statement == "C":
        headers, rows = scrap_table(html)
    else:
        _, headers, rows = source_table(html, required)
    indices = {key(label): i for i, label in enumerate(headers)}
    wanted_month = month_end(period).strftime("%b-%Y").upper()
    chosen = [row for row in rows
              if norm(row[indices["DEPOT"]]) == norm(depot_code)
              and norm(row[indices["MONTHYEAR"]]) == wanted_month
              and norm(row[indices["TYRESIZE"]]) == "ALL TYRE SIZES TOTAL"]
    unique = []
    for row in chosen:
        if row not in unique:
            unique.append(row)
    if len(unique) != 1:
        raise ValueError(f"{statement} {period}: no unique exact depot/month/total-size row")
    identifiers = {"MONTHYEAR", "ZONENAME", "REGION", "DEPOT", "DCPCODE", "TYRESIZE"}
    typed = [text if key(label) in identifiers else number(text)
             for label, text in zip(headers, unique[0])]
    return dict(depot=norm(depot), period=period, group=statement,
                headers=headers, rows=[typed], tyre_code=depot_code,
                through=month_end(period).isoformat(),
                source=f"{BASE}/tyres/{statement.lower()}_statement_final.php")


def parse_engine(html, depot, period, mode, through=None):
    if mode not in ("UD", "UM"):
        raise ValueError("Engine request must explicitly select UD or UM")
    soup, headers, rows = source_table(html, ["SNO", "ENGINE MODEL", "Grand Total"])
    heading = next((h.get_text(" ", strip=True) for h in soup.find_all(re.compile("^h[1-6]$"))
                    if "ENGINE MODEL" in norm(h.get_text(" ", strip=True))), "")
    expected_end = through or month_end(period)
    expected_start = expected_end.replace(day=1) if mode == "UD" else fy_start(period)
    dates = re.search(r"FROM:\s*(\d{2}[-/]\d{2}[-/]\d{4})\s+TO\s+(\d{2}[-/]\d{2}[-/]\d{4})",
                      norm(heading))
    prefix = norm(heading).split(" DEPOT", 1)[0]
    if prefix != norm(depot) or not dates:
        raise ValueError("Engine report depot/period heading is not verified")
    actual = [datetime.strptime(x.replace("/", "-"), "%d-%m-%Y").date() for x in dates.groups()]
    if actual != [expected_start, expected_end]:
        raise ValueError(f"{mode} returned a different period: {actual}")
    typed, identities, totals = [], set(), 0
    for row in rows:
        is_total = key(row[0]) == "TOTAL"
        identity = "TOTAL" if is_total else norm(row[1])
        if not identity or identity in identities:
            raise ValueError("Engine source contains duplicate/blank row identities")
        identities.add(identity)
        if is_total:
            totals += 1
            if row[1]:
                raise ValueError("TOTAL label colspan must cover SNO and ENGINE MODEL")
        elif not re.fullmatch(r"\d+", row[0]):
            raise ValueError("Engine row serial is not recognised")
        typed.append([row[0] if is_total else int(row[0]), row[1]] +
                     [number(value) for value in row[2:]])
    if totals != 1 or len(typed) < 2 or key(typed[-1][0]) != "TOTAL":
        raise ValueError("Engine matrix requires engine rows and one final source TOTAL")
    return dict(depot=norm(depot), period=period, group=mode, headers=headers, rows=typed,
                start=expected_start.isoformat(), through=expected_end.isoformat(),
                heading=heading, source=f"{BASE}/mednew/eng_prod_um.php")


class SourceAdapter:
    def __init__(self, session, depot, region, tyre_route, today=None):
        self.session, self.depot, self.region = session, depot, region
        self.tyre_route, self.today = tyre_route, today

    def __call__(self, group, period):
        year, month = map(int, period.split("-"))
        if group in STATEMENTS:
            route = self.tyre_route(self.depot, year, month)
            code = route[0] if isinstance(route, (tuple, list)) else route
            payload = dict(zone="", region="", depot=code,
                           month_year=month_end(period).strftime("%b-%Y").upper(),
                           tyre_size="All Tyre Sizes Total")
            url = f"{BASE}/tyres/{group.lower()}_statement_final.php"
            response = self.session.get(url, params=payload, timeout=45)
            response.raise_for_status()
            snapshot = parse_tyre(response.text, group, code, period, self.depot)
        else:
            last = effective_end(period, self.today)
            payload = dict(fdate=f"{last.day}/{last.month}/{last.year}", reg=self.region,
                           dept=self.depot, fupto=group)
            response = self.session.post(f"{BASE}/mednew/eng_prod_um.php", data=payload, timeout=45)
            response.raise_for_status()
            snapshot = parse_engine(response.text, self.depot, period, group, last)
        snapshot["request"] = payload
        snapshot["fetched_at"] = datetime.now(IST).isoformat(timespec="seconds")
        snapshot["provisional"] = snapshot["through"] != month_end(period).isoformat()
        # Tyre statements for an open month can still be changing, even when their row has only a month key.
        if group in STATEMENTS:
            today = self.today or datetime.now(IST).date()
            snapshot["provisional"] = month_end(period) >= today
        return snapshot


def new_cache(depot):
    return dict(schema=SCHEMA, depot=norm(depot), months={}, errors={})


def validate_snapshot(snapshot, depot, group, period):
    if not isinstance(snapshot, dict) or snapshot.get("depot") != norm(depot):
        raise ValueError("Detail snapshot belongs to another depot")
    if snapshot.get("period") != period or snapshot.get("group") != group:
        raise ValueError("Detail snapshot belongs to another source period/mode")
    headers, rows = snapshot.get("headers", []), snapshot.get("rows", [])
    if not headers or not rows or any(len(row) != len(headers) for row in rows):
        raise ValueError("Incomplete detail snapshot")
    if len(set(map(key, headers))) != len(headers):
        raise ValueError("Duplicate cached source columns")
    for row in rows:
        for value in row:
            if isinstance(value, bool) or not isinstance(value, (str, int, float, type(None))):
                raise ValueError("Invalid detail value")
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError("Non-finite detail value")
    return snapshot


def encode(cache):
    raw = json.dumps(cache, ensure_ascii=True, allow_nan=False, separators=(",", ":"))
    chunks = [raw[i:i+30000] for i in range(0, len(raw), 30000)]
    return [[SCHEMA, cache["depot"], len(chunks), hashlib.sha256(raw.encode()).hexdigest()]] + [[c] for c in chunks]


def decode(rows, depot):
    if not rows or len(rows[0]) < 4 or list(rows[0][:2]) != [SCHEMA, norm(depot)]:
        raise ValueError("Detail history identity/schema mismatch; refusing to replace saved history")
    count = int(rows[0][2])
    if count < 1 or len(rows) < count + 1 or any(not r or not isinstance(r[0], str) for r in rows[1:count+1]):
        raise ValueError("Detail history chunks are incomplete")
    raw = "".join(r[0] for r in rows[1:count+1])
    if hashlib.sha256(raw.encode()).hexdigest() != rows[0][3]:
        raise ValueError("Detail history checksum mismatch")
    cache = json.loads(raw)
    if cache.get("schema") != SCHEMA or cache.get("depot") != norm(depot):
        raise ValueError("Detail history identity mismatch")
    for period, groups in cache["months"].items():
        month_end(period)
        for group, snapshot in groups.items():
            if group not in ("B", "C", "F", "UD", "UM"):
                raise ValueError("Unknown cached detail source")
            validate_snapshot(snapshot, depot, group, period)
    cache.setdefault("errors", {})
    return cache


def update(cache, wanted, fetch, selected_um=None, checkpoint=lambda: None, refresh_monthly=False):
    """History is immutable after a verified closed-month fetch; UM is a dated refresh."""
    calls = []
    for period in wanted:
        if period < FIRST_PERIOD and not refresh_monthly:
            continue
        month_end(period)
        groups = cache["months"].setdefault(period, {})
        for group in ("B", "C", "F", "UD"):
            old = groups.get(group)
            if old and not old.get("provisional") and not refresh_monthly:
                continue
            _fetch_into(cache, group, period, fetch, calls)
        checkpoint()
    if selected_um and selected_um >= FIRST_PERIOD:
        _fetch_into(cache, "UM", selected_um, fetch, calls)
        checkpoint()
    return calls


def _fetch_into(cache, group, period, fetch, calls):
    calls.append((group, period))
    try:
        fresh = validate_snapshot(fetch(group, period), cache["depot"], group, period)
        # Replace the whole verified source matrix, so rows absent on a selected
        # refresh cannot survive as stale cells. Failures never erase the old matrix.
        cache["months"].setdefault(period, {})[group] = copy.deepcopy(fresh)
        cache["errors"].get(period, {}).pop(group, None)
    except Exception as exc:
        message = f"{type(exc).__name__}: {exc}"[:300]
        cache["errors"].setdefault(period, {})[group] = message
        print(f"DETAIL_SOURCE_UNAVAILABLE {cache['depot']} {period} {group}: saved values retained; {message}")


def read_google_cache(service, spreadsheet_id, depot):
    api = service.spreadsheets()
    meta = api.get(spreadsheetId=spreadsheet_id, fields="sheets.properties").execute()
    if not any(s["properties"]["title"] == CACHE_TITLE for s in meta.get("sheets", [])):
        return new_cache(depot)
    rows = api.values().get(spreadsheetId=spreadsheet_id, range=f"'{CACHE_TITLE}'!A:D").execute().get("values", [])
    return decode(rows, depot)


def save_google_cache(service, spreadsheet_id, cache):
    api = service.spreadsheets()
    meta = api.get(spreadsheetId=spreadsheet_id, fields="sheets.properties").execute()
    existing = next((s for s in meta.get("sheets", []) if s["properties"]["title"] == CACHE_TITLE), None)
    rows = encode(cache)
    if existing is None:
        result = api.batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": [{"addSheet": {
            "properties": {"title": CACHE_TITLE, "hidden": True,
                           "gridProperties": {"rowCount": len(rows) + 2, "columnCount": 4}}}}]}).execute()
        sid = result["replies"][0]["addSheet"]["properties"]["sheetId"]
    else:
        sid = existing["properties"]["sheetId"]
        # Do not shrink until the checksum-protected replacement is saved.
        count = max(existing["properties"].get("gridProperties", {}).get("rowCount", 0), len(rows) + 2)
        api.batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": [{"updateSheetProperties": {
            "properties": {"sheetId": sid, "hidden": True, "gridProperties": {"rowCount": count, "columnCount": 4}},
            "fields": "hidden,gridProperties.rowCount,gridProperties.columnCount"}}]}).execute()
    api.values().update(spreadsheetId=spreadsheet_id, range=f"'{CACHE_TITLE}'!A1",
                        valueInputOption="RAW", body={"values": rows}).execute()


def sections(cache, selected, annual):
    # The FY2025-26 floor applies to annual history. A standalone Monthly
    # selection may request any supported source month without backfilling FYs.
    periods = annual_periods(selected) if annual else [selected]
    from report_tyre_template import tyre_sections
    tyre = tyre_sections(cache, periods, annual)
    engine = []
    if annual and selected >= FIRST_PERIOD:
        snap = cache["months"].get(selected, {}).get("UM")
        engine.append(_engine_section(cache, selected, "UM", snap, selected=True))
    for period in periods:
        snap = cache["months"].get(period, {}).get("UD")
        engine.append(_engine_section(cache, period, "UD", snap))
    if not periods:
        for target in (tyre, engine):
            target.append(dict(kind="note", title="Source history starts in FY2025-26",
                               headers=[], rows=[], missing=[], stale=[]))
    return {TYRE_TITLE: tyre, ENGINE_TITLE: engine}


def _engine_section(cache, period, mode, snap, selected=False):
    label = month_end(period).strftime("%B %Y")
    title = f"FY {fy_label(period)} · {label} · " + ("UM — selected cumulative" if mode == "UM" else "UD — monthly actual")
    return dict(kind="engine", title=title, headers=snap["headers"] if snap else [],
                rows=snap["rows"] if snap else [], missing=[] if snap else [period],
                stale=[period] if cache["errors"].get(period, {}).get(mode) and snap else [],
                source=f"{BASE}/mednew/eng_prod_um.php", selected=selected,
                period_text=f"{snap['start']} to {snap['through']}" if snap else "",
                provisional=bool(snap and snap.get("provisional")))


def render_tabs(workbook, cache, selected, annual=False):
    """Rebuild only the two generated views. Cache and unrelated sheets survive."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    from report_branding import add_logo, style_kmpl
    from report_tyre_template import write_inputs, finish_totals
    write_inputs(workbook, cache, annual_periods(selected) if annual else [selected])
    result = {}
    for title, blocks in sections(cache, selected, annual).items():
        index = workbook.sheetnames.index(title) if title in workbook.sheetnames else len(workbook.sheetnames)
        if title in workbook.sheetnames:
            del workbook[title]
        ws = workbook.create_sheet(title, index)
        # Seven columns fit the identity/logo and the smallest verified source
        # matrix. Do not retain July-only product columns in a May monthly view.
        width = max([len(b["headers"]) for b in blocks] + [7])
        last = get_column_letter(width)
        heading_last = "N" if title == TYRE_TITLE else last
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = BLUE if title == TYRE_TITLE else "287B59"
        ws.merge_cells(f"A2:{heading_last}2")
        ws["A2"] = f"APSRTC — {cache['depot']} DEPOT — {title.upper()}"
        ws["A2"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 9
        ws.row_dimensions[2].height = 27
        ws.merge_cells(f"A3:C3")
        ws.merge_cells(f"D3:{heading_last}3")
        ws["D3"] = ("ANNUAL KPI · Through " if annual else "MONTHLY REPORT · ") + month_end(selected).strftime("%B %Y")
        ws["D3"].font = Font(name="Arial", size=12, bold=True, color=BLUE)
        ws["D3"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"A4:{heading_last}4")
        ws["A4"] = (("All Tyre Sizes Total · FY totals use source counts; percentages use matching denominators" if annual else
                     "All Tyre Sizes Total · Source monthly percentages · blank: no source value") if title == TYRE_TITLE else
                    ("UD: monthly actual · UM: April through selected month · blank: no source value" if annual else
                     "UD: monthly actual through month end · blank: no source value"))
        ws["A4"].font = Font(name="Arial", size=10, italic=True, color="63758A")
        ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 22
        r, first_header, engine_blocks = 6, None, 0
        for block in blocks:
            columns = max(len(block["headers"]), 3)
            end = get_column_letter(columns)
            ws.merge_cells(f"A{r}:{end}{r}")
            ws.cell(r, 1, block["title"])
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="287B59" if block.get("selected") else NAVY)
            ws.cell(r, 1).font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            ws.cell(r, 1).alignment = Alignment(vertical="center")
            ws.row_dimensions[r].height = 27
            r += 1
            context = block.get("period_text") or "Month-wise source values"
            if block.get("provisional"):
                context += (" · PROVISIONAL — current monthly statement; no daily cutoff supplied"
                            if block["kind"] == "tyre" else " · PROVISIONAL — through completed days")
            ws.merge_cells(f"A{r}:{end}{r}")
            ws.cell(r, 1, context)
            ws.cell(r, 1).font = Font(name="Arial", size=10, color="63758A")
            ws.row_dimensions[r].height = 22
            r += 1
            if block["headers"]:
                if block.get("header_groups"):
                    for start, stop, label in block["header_groups"]:
                        if start != stop:
                            ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=stop)
                        cell = ws.cell(r, start, label)
                        cell.fill = PatternFill("solid", fgColor=BLUE)
                        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    ws.row_dimensions[r].height = 29
                    r += 1
                first_header = first_header or r
                for c, label in enumerate(block["headers"], 1):
                    cell = ws.cell(r, c, label)
                    cell.fill = PatternFill("solid", fgColor=BLUE)
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = Border(right=Side(style="thin", color="FFFFFF"))
                ws.row_dimensions[r].height = 58 if title == TYRE_TITLE else 40
                r += 1
                data_start = r
                for row_index, row in enumerate(block["rows"]):
                    total = (block["kind"] == "engine" and key(row[0]) == "TOTAL") or row_index == block.get("total_index")
                    for c, value in enumerate(row, 1):
                        cell = ws.cell(r, c, value)
                        if isinstance(value, str) and value.startswith("="):
                            cell.data_type = "s"
                        cell.fill = PatternFill("solid", fgColor="EAF2FA" if total else (PALE if r % 2 else "FFFFFF"))
                        cell.font = Font(name="Arial", size=10, color=NAVY if total else INK, bold=total)
                        cell.alignment = Alignment(horizontal="right" if isinstance(value, (int, float)) else "left",
                                                   vertical="center", wrap_text=c in (3, 4, 7, 8) and title == TYRE_TITLE)
                        cell.border = Border(bottom=Side(style="thin", color=LINE))
                        label = block["headers"][c-1]
                        if isinstance(value, (int, float)):
                            cell.number_format = "0.00%" if c in block.get("percent_columns", []) else ("0.00" if block["kind"] == "engine" and c > 2 or "%" in label else "#,##0")
                            if block["kind"] == "engine" and c > 2:
                                style_kmpl(cell)
                                if total:
                                    cell.font = Font(name="Arial", size=10, color=cell.font.color, bold=True)
                    if total and block["kind"] == "engine":
                        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
                    ws.row_dimensions[r].height = 26 if title == TYRE_TITLE else 23
                    r += 1
                finish_totals(ws, block, data_start)
            for note in block.get("notes", []):
                ws.merge_cells(f"A{r}:N{r}")
                ws.cell(r, 1, note)
                ws.cell(r, 1).font = Font(name="Arial", size=10, color="9C6500")
                ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[r].height = 32
                r += 1
            for label, periods in (("Source unavailable", block["missing"]), ("Saved values retained; latest refresh failed", block["stale"])):
                if periods:
                    ws.merge_cells(f"A{r}:{end}{r}")
                    ws.cell(r, 1, label + ": " + ", ".join(periods))
                    ws.cell(r, 1).font = Font(name="Arial", size=10, color="9C6500", bold=True)
                    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
                    ws.row_dimensions[r].height = max(23, 18 * math.ceil(len(ws.cell(r, 1).value) / max(70, width*8)))
                    r += 1
            if block.get("source"):
                ws.merge_cells(f"A{r}:{end}{r}")
                ws.cell(r, 1, "Source: " + block["source"])
                ws.cell(r, 1).font = Font(name="Arial", size=9, italic=True, color="63758A")
                ws.row_dimensions[r].height = 19
                r += 1
            if block["kind"] == "engine":
                engine_blocks += 1
                if annual and engine_blocks % 2 == 0:
                    ws.row_breaks.append(Break(id=r))
            r += 2
        widths = ([6, 13] if title == TYRE_TITLE else [6, 25])
        for c in range(1, width + 1):
            ws.column_dimensions[get_column_letter(c)].width = widths[c-1] if c <= len(widths) else 12
        # A3:C3 is reserved for the embedded logo; size it to that exact span.
        logo_width = sum(ws.column_dimensions[get_column_letter(c)].width * 7 + 5 for c in (1, 2, 3)) - 10
        add_logo(ws, anchor="A3", width=logo_width)
        ws.row_dimensions[3].height = math.ceil(logo_width * 130 / 640 * .75) + 3
        ws.freeze_panes = f"C{(first_header or 7) + 1}"
        ws.print_title_rows = "1:4"
        ws.print_area = f"A1:{last}{max(4, r - 2)}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = f"APSRTC | {cache['depot']} | {month_end(selected).strftime('%B %Y')}"
        ws.oddFooter.right.text = "Page &P of &N"
        result[title] = ws
    return result


def write_xlsx_cache(workbook, cache):
    if CACHE_TITLE in workbook.sheetnames:
        del workbook[CACHE_TITLE]
    ws = workbook.create_sheet(CACHE_TITLE)
    for row in encode(cache):
        ws.append(row)
    ws.sheet_state = "hidden"


def _rgb(value):
    return {component: int(value[i:i+2], 16) / 255 for component, i in (("red", 0), ("green", 2), ("blue", 4))}


def google_requests(worksheet, sheet_id, existing=None):
    """Replace the owned view's complete grid, formats and merges, then shrink it."""
    from openpyxl.utils import get_column_letter
    content_rows, cols = worksheet.max_row, worksheet.max_column
    rows = content_rows + 2
    requests = [
        {"unmergeCells": {"range": {"sheetId": sheet_id}}},
        {"clearBasicFilter": {"sheetId": sheet_id}},
        {"updateSheetProperties": {"properties": {"sheetId": sheet_id, "gridProperties": {
            "frozenRowCount": 0, "frozenColumnCount": 0}}, "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}},
    ]
    for i in reversed(range(len((existing or {}).get("conditionalFormats", [])))):
        requests.append({"deleteConditionalFormatRule": {"sheetId": sheet_id, "index": i}})
    for band in (existing or {}).get("bandedRanges", []):
        requests.append({"deleteBanding": {"bandedRangeId": band["bandedRangeId"]}})
    for chart in (existing or {}).get("charts", []):
        requests.append({"deleteEmbeddedObject": {"objectId": chart["chartId"]}})
    requests.append({"updateSheetProperties": {"properties": {"sheetId": sheet_id, "hidden": worksheet.title == TYRE_INPUT_TITLE,
        "gridProperties": {"rowCount": rows, "columnCount": cols, "hideGridlines": True},
        "tabColor": _rgb(BLUE if worksheet.title == TYRE_TITLE else "287B59")},
        "fields": "hidden,gridProperties.rowCount,gridProperties.columnCount,gridProperties.hideGridlines,tabColor"}})
    body = []
    for r in range(1, rows + 1):
        values = []
        for c in range(1, cols + 1):
            if r > content_rows:
                values.append({})
                continue
            cell = worksheet.cell(r, c)
            value = cell.value
            fmt = dict(horizontalAlignment=(cell.alignment.horizontal or "left").upper(),
                       verticalAlignment="MIDDLE", wrapStrategy="WRAP" if cell.alignment.wrap_text else "CLIP",
                       textFormat=dict(fontFamily="Arial", fontSize=int(cell.font.sz or 10), bold=bool(cell.font.bold),
                                       italic=bool(cell.font.italic)))
            color = cell.font.color
            if color and color.type == "rgb":
                fmt["textFormat"]["foregroundColor"] = _rgb(color.rgb[-6:])
            if cell.fill.patternType == "solid":
                fmt["backgroundColor"] = _rgb(cell.fill.fgColor.rgb[-6:])
            if cell.number_format not in ("General", "@"):
                fmt["numberFormat"] = dict(type="PERCENT" if "%" in cell.number_format else "NUMBER", pattern=cell.number_format)
            if cell.border.bottom is not None and cell.border.bottom.style:
                fmt["borders"] = {"bottom": {"style": "SOLID", "color": _rgb(LINE)}}
            record = {"userEnteredFormat": fmt}
            if cell.data_type == "f":
                record["userEnteredValue"] = {"formulaValue": value}
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                record["userEnteredValue"] = {"numberValue": value}
            elif value is not None:
                # Source strings are literal values, even when they start with '='.
                record["userEnteredValue"] = {"stringValue": str(value)}
            elif r == 3 and c == 1 and worksheet.title != TYRE_INPUT_TITLE:
                # Sheets v4 cannot create an over-grid image. Use native text
                # branding in the reserved logo span; XLSX contains the logo.
                record["userEnteredValue"] = {"stringValue": "APSRTC"}
                record["userEnteredFormat"]["textFormat"] = {
                    "fontFamily": "Arial", "fontSize": 22, "bold": True,
                    "foregroundColor": _rgb(NAVY)}
                record["userEnteredFormat"]["horizontalAlignment"] = "CENTER"
            values.append(record)
        body.append({"values": values})
    requests.append({"updateCells": {"range": {"sheetId": sheet_id, "endRowIndex": rows, "endColumnIndex": cols},
                     "rows": body, "fields": "userEnteredValue,userEnteredFormat,note"}})
    for merged in worksheet.merged_cells.ranges:
        requests.append({"mergeCells": {"range": {"sheetId": sheet_id, "startRowIndex": merged.min_row-1,
            "endRowIndex": merged.max_row, "startColumnIndex": merged.min_col-1, "endColumnIndex": merged.max_col},
            "mergeType": "MERGE_ALL"}})
    for c in range(1, cols + 1):
        pixels = round(worksheet.column_dimensions[get_column_letter(c)].width * 7 + 5)
        requests.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS",
            "startIndex": c-1, "endIndex": c}, "properties": {"pixelSize": pixels}, "fields": "pixelSize"}})
    for r in range(1, rows + 1):
        pixels = round((worksheet.row_dimensions[r].height or 22) * 4/3)
        requests.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "ROWS",
            "startIndex": r-1, "endIndex": r}, "properties": {"pixelSize": pixels}, "fields": "pixelSize"}})
    # Only the first four identity rows are frozen: freezing through the first
    # section merge would cross a vertical boundary when source widths change.
    requests.append({"updateSheetProperties": {"properties": {"sheetId": sheet_id, "gridProperties": {
        "frozenRowCount": 0 if worksheet.title == TYRE_INPUT_TITLE else 4, "frozenColumnCount": 0}}, "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}})
    return requests


def publish_google_tabs(service, spreadsheet_id, workbook):
    api = service.spreadsheets()
    for title in (TYRE_INPUT_TITLE, *TAB_TITLES):
        meta = api.get(spreadsheetId=spreadsheet_id).execute()
        existing = next((s for s in meta.get("sheets", []) if s["properties"]["title"] == title), None)
        if existing is None:
            result = api.batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": [{"addSheet": {
                "properties": {"title": title}}}]}).execute()
            sid = result["replies"][0]["addSheet"]["properties"]["sheetId"]
        else:
            sid = existing["properties"]["sheetId"]
        api.batchUpdate(spreadsheetId=spreadsheet_id, body={"requests":
            google_requests(workbook[title], sid, existing)}).execute()


def attach_monthly_tabs(workbook, session, depot, region, selected, folder):
    """Load only the same depot/month's detail cache, preserving existing vehicle seeding."""
    from annual_kpi_runner_v3 import tyre_site_info
    from src.integrations.google_drive import find_file
    from src.integrations.google_sheets import sheets_service
    existing = find_file(folder, f"{depot}_{selected}")
    cache = read_google_cache(sheets_service(), existing["id"], depot) if existing else new_cache(depot)
    adapter = SourceAdapter(session, depot, region, tyre_site_info)
    update(cache, [selected], adapter, refresh_monthly=True)
    render_tabs(workbook, cache, selected)
    write_xlsx_cache(workbook, cache)
    return cache
