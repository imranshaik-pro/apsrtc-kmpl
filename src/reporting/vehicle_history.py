"""Vehicle Performance History sheet for Monthly reports."""
from __future__ import annotations
from datetime import datetime
from io import StringIO
import re
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MTD_URL = "http://103.44.14.20/med/edeengine.php"
TREND_URL = "http://103.44.14.20/med/vkmpl12trend.php"
S3_URL = "http://103.44.14.20/med/s3cover.php"
S4_URL = "http://103.44.14.20/med/s4cover.php"
FYS = ("2024-25", "2025-26", "2026-27")
MONTHS = ("Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar")
THIN = Side(style="thin", color="B7C9E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def norm_vehicle(value):
    """Canonical APSRTC vehicle number: uppercase, no spaces, no leading AP."""
    vehicle = re.sub(r"\s+", "", str(value or "")).upper()
    return vehicle[2:] if vehicle.startswith("AP") else vehicle

def _clean_columns(df):
    if isinstance(df.columns, pd.MultiIndex):
        cols=[]
        for col in df.columns:
            parts=[str(x).strip() for x in col if str(x).strip() and not str(x).startswith("Unnamed")]
            cols.append(" ".join(dict.fromkeys(parts)))
        df.columns=cols
    else: df.columns=[str(x).strip() for x in df.columns]
    return df

def _table(html, required):
    for df in pd.read_html(StringIO(html)):
        df=_clean_columns(df); joined=" | ".join(map(str,df.columns)).lower()
        if all(x.lower() in joined for x in required): return df
    return None

def fetch_mtd(session, yyyymm, regn, depot):
    payload={"yymm":yyyymm,"regn":regn,"depot":depot,"stype":"","eng":"","kms":"","kmsl":"","kmpl":"","kmpll":"","fstatus":"0","veh":""}
    r=session.post(MTD_URL,data=payload,timeout=30); r.raise_for_status()
    df=_table(r.text,["Veh No","HSD KMPL","Comm Date"])
    if df is None: raise RuntimeError(f"MTD-598 table not found for {depot} {yyyymm}")
    result={}
    for _,row in df.iterrows():
        def get(label):
            for c in df.columns:
                if label.lower()==str(c).strip().lower(): return row[c]
            return ""
        v=norm_vehicle(get("Veh No"))
        if not v or v=="NAN": continue
        try: kmpl=float(get("HSD KMPL"))
        except (TypeError,ValueError): kmpl=None
        result[v]={"kmpl":kmpl,"op":str(get("Veh Type")).strip(),"engine":str(get("Eng Make")).strip(),"comm_date":str(get("Comm Date")).strip()}
    return result

def fetch_trend(session, fy, zone, regn, depot):
    fyymm={"2025-26":"202603March_2026","2026-27":"202703March_2027"}[fy]
    r=session.post(TREND_URL,data={"yymm":"","fyymm":fyymm,"zone":zone,"regn":regn,"dept":depot,"submit":"SUBMIT"},timeout=30); r.raise_for_status()
    df=_table(r.text,["Vehicle No","Product Type","Engine Type"])
    if df is None: raise RuntimeError(f"12-month trend table not found for {depot} {fy}")
    result={}
    for _,row in df.iterrows():
        def find(prefix,contains=None):
            for c in df.columns:
                s=str(c)
                if s.lower().startswith(prefix.lower()) and (contains is None or contains.lower() in s.lower()): return row[c]
            return ""
        v=norm_vehicle(find("Vehicle No"))
        if not v or v=="NAN": continue
        months={}; start_year=int(fy[:4])
        for i,mon in enumerate(MONTHS):
            yy=start_year if i<9 else start_year+1; label=f"{mon}-{str(yy)[2:]}"; val=find(label,"HSD KMPL")
            try: months[mon]=float(val)
            except (TypeError,ValueError): months[mon]=None
        result[v]={"op":str(find("Product Type")).strip(),"engine":str(find("Engine Type")).strip(),"months":months}
    return result

def fetch_schedule(session,schedule_no,yyyymm,month_label,zone,regn,depot):
    url=S3_URL if schedule_no==3 else S4_URL
    r=session.post(url,data={"yymm":f"{yyyymm}{month_label}","zone":zone,"regn":regn,"dept":depot,"submit":"SUBMIT"},timeout=30); r.raise_for_status()
    date_col=f"Date of Sch-{'III' if schedule_no==3 else 'IV'}"; df=_table(r.text,["Vehicle No",date_col])
    if df is None: return {}
    events={}
    for _,row in df.iterrows():
        vcol=next((c for c in df.columns if str(c).strip().lower()=="vehicle no."),None) or next((c for c in df.columns if "vehicle no" in str(c).lower()),None)
        dcol=next((c for c in df.columns if date_col.lower() in str(c).lower()),None)
        if vcol is None or dcol is None: continue
        v=norm_vehicle(row[vcol]); d=str(row[dcol]).strip()
        if not v or v=="NAN" or not d or d.lower()=="nan": continue
        events.setdefault(v,[]).append(d)
    return events

def _month_pairs(fy):
    y=int(fy[:4]); return [(MONTHS[i],f"{y if i<9 else y+1}{i+4:02d}" if i<9 else f"{y+1}{i-8:02d}") for i in range(12)]

def _selected_month_index(year,month): return month-4 if month>=4 else month+8

def read_existing_history(ws):
    out={}
    if ws is None:return out
    headers={str(c.value).strip():c.column for c in ws[1] if c.value is not None}
    if "Veh No" not in headers or "FY Year" not in headers:return out
    last_vehicle=""
    for r in range(2,ws.max_row+1):
        raw_v=ws.cell(r,headers["Veh No"]).value
        if raw_v not in (None,""):last_vehicle=norm_vehicle(raw_v)
        fy=str(ws.cell(r,headers["FY Year"]).value or "").strip()
        if not last_vehicle or fy not in FYS:continue
        out.setdefault(last_vehicle,{})[fy]={m:ws.cell(r,headers[m]).value for m in MONTHS if m in headers}
    return out

def build_history(session,selected_year,selected_month,zone,regn,depot,existing=None,provisional_roster=None):
    selected_yyyymm=f"{selected_year}{selected_month:02d}"
    roster=fetch_mtd(session,selected_yyyymm,regn,depot)
    # Open-month reports may not yet have an official MTD-598 population.
    # In that case the caller supplies the operational roster observed in daily data/events.
    if not roster and provisional_roster:
        roster={v:dict(info) for v,info in provisional_roster.items()}
    existing=existing or {}
    values={v:{fy:dict(existing.get(v,{}).get(fy,{})) for fy in FYS} for v in roster}; new_vehicles={v for v in roster if v not in existing}; need_full_init=not existing
    targets_2425=set(roster) if need_full_init else new_vehicles
    if targets_2425:
        for mon,ym in _month_pairs("2024-25"):
            snap=fetch_mtd(session,ym,regn,depot)
            for v in targets_2425:
                if v in snap:values[v]["2024-25"][mon]=snap[v]["kmpl"]
    for fy in ("2025-26","2026-27"):
        targets=set(roster) if need_full_init else new_vehicles
        if fy=="2026-27" and (selected_year,selected_month)>=(2026,4):targets|=set(roster)
        if not targets:continue
        trend=fetch_trend(session,fy,zone,regn,depot)
        for v in targets:
            if v not in trend:continue
            for mon,val in trend[v]["months"].items():
                if fy=="2026-27":
                    mi=MONTHS.index(mon)
                    if (selected_year,selected_month)<(2027,4) and mi>_selected_month_index(selected_year,selected_month):continue
                if need_full_init or v in new_vehicles or mon==MONTHS[_selected_month_index(selected_year,selected_month)]:values[v][fy][mon]=val
            if not roster[v]["op"] and trend[v]["op"]:roster[v]["op"]=trend[v]["op"]
            if not roster[v]["engine"] and trend[v]["engine"]:roster[v]["engine"]=trend[v]["engine"]
    remarks=[]
    if (selected_year,selected_month)>=(2026,4):
        end_idx=min(_selected_month_index(selected_year,selected_month),11); indices=range(0,end_idx+1) if need_full_init or new_vehicles else [end_idx]
        for i in indices:
            mon=MONTHS[i]; yy=2026 if i<9 else 2027; mm=i+4 if i<9 else i-8; ym=f"{yy}{mm:02d}"; month_label=datetime(yy,mm,1).strftime("%B_%Y")
            s3=fetch_schedule(session,3,ym,month_label,zone,regn,depot); s4=fetch_schedule(session,4,ym,month_label,zone,regn,depot)
            targets=set(roster) if need_full_init else (new_vehicles|(set(roster) if i==end_idx else set()))
            for v in targets:
                e3,e4=s3.get(v,[]),s4.get(v,[]); abnormal=len(e3)>1 or len(e4)>1 or (e3 and e4)
                if abnormal:
                    details=[]
                    if e3:details.append("SCH-III="+", ".join(e3))
                    if e4:details.append("SCH-IV="+", ".join(e4))
                    remarks.append((f"{mon}-{str(yy)[2:]}",v,"; ".join(details)));continue
                event=(3,e3[0]) if e3 else ((4,e4[0]) if e4 else None)
                if event:
                    try:day=int(event[1].split("-")[0])
                    except Exception:day=event[1]
                    base=values[v]["2026-27"].get(mon); base_text="" if base in (None,"") else (f"{base:.2f}" if isinstance(base,(int,float)) else str(base).split(" ")[0]); icon="🔧" if event[0]==3 else "⚙"
                    values[v]["2026-27"][mon]=f"{base_text} {icon} {day}".strip()
    return roster,values,remarks

def write_vehicle_360_sheet(wb, roster, values, events=None):
    """Create a management-friendly per-vehicle 360 summary.

    Manual event data is optional until the live Event Register is connected.
    APSRTC-derived KMPL identity remains available immediately.
    """
    from .vehicle_events import aggregate_change_dates

    if "Vehicle 360" in wb.sheetnames:
        del wb["Vehicle 360"]
    ws = wb.create_sheet("Vehicle 360")
    ws.sheet_view.showGridLines = False
    ws.append(["APSRTC – VEHICLE 360° HISTORY"])
    ws.append(["KMPL performance + maintenance + aggregate / breakdown / tyre event history"])
    ws.append([])
    headers = ["S.NO","Veh No","OP Type","Eng Type","Latest KMPL","Major Aggregate Changes","Breakdowns","Tyre Changes"]
    ws.append(headers)
    ws.merge_cells("A1:H1"); ws.merge_cells("A2:H2")
    ws["A1"].font=Font(bold=True,color="FFFFFF",size=16); ws["A1"].fill=PatternFill("solid",fgColor="17365D"); ws["A1"].alignment=Alignment(horizontal="center",vertical="center")
    ws["A2"].font=Font(bold=True,color="1F4E78",size=10); ws["A2"].alignment=Alignment(horizontal="center",vertical="center")
    for c in ws[4]:
        c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=BORDER
    by_vehicle={}
    for e in events or []:
        by_vehicle.setdefault(e.vehicle_no,[]).append(e)
    ordered=sorted(roster,key=lambda v:(str(roster[v].get("op","")).upper(),str(roster[v].get("engine","")).upper(),v))
    for idx,v in enumerate(ordered,1):
        ve=by_vehicle.get(v,[])
        changes=aggregate_change_dates(ve)
        aggregate_text="\n".join(f"{component}: {' / '.join(dates)}" for component,dates in changes.items())
        breakdown_text="\n".join(
            f"{e.event_date}: " +
            " | ".join(x for x in (
                e.breakdown_location,
                (f"{e.kms_cancelled} KM cancelled" if e.kms_cancelled else ""),
                e.breakdown_details,
                e.remarks,
            ) if x)
            for e in ve if e.event_type=="BREAKDOWN"
        )
        tyre_text="\n".join(
            f"{e.event_date}: " + " | ".join(x for x in (e.tyre_history, e.remarks) if x)
            for e in ve if e.event_type=="TYRE CHANGE"
        )
        latest=""
        for fy in reversed(FYS):
            for mon in reversed(MONTHS):
                val=values.get(v,{}).get(fy,{}).get(mon,"")
                if val not in (None,""):
                    latest=str(val).split(" ")[0]; break
            if latest: break
        ws.append([idx,v,roster[v].get("op",""),roster[v].get("engine",""),latest,aggregate_text,breakdown_text,tyre_text])
        r=ws.max_row
        for c in range(1,9):
            ws.cell(r,c).border=BORDER; ws.cell(r,c).alignment=Alignment(horizontal="center" if c in (1,2,5) else "left",vertical="top",wrap_text=True)
        ws.row_dimensions[r].height=36
    for i,w in enumerate([7,15,18,18,13,38,38,38],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="E5"; ws.auto_filter.ref=f"A4:H{ws.max_row}"
    return ws


def write_history_sheet(wb,roster,values,remarks,population_note=None):
    if "Vehicle Performance" in wb.sheetnames:del wb["Vehicle Performance"]
    ws=wb.create_sheet("Vehicle Performance"); headers=["S.NO","Veh No","OP Type","Eng Type","FY Year",*MONTHS]
    ws.append(["APSRTC – VEHICLE PERFORMANCE HISTORY"])
    ws.append([population_note or "3 Financial Year HSD KMPL History | Current vehicle population based on selected-month MTD-598"])
    ws.append(["🔧 Schedule-III completed   |   ⚙ Schedule-IV completed   |   Number = completion day"])
    ws.append([])
    ws.append(headers)
    ws.merge_cells("A1:Q1"); ws.merge_cells("A2:Q2"); ws.merge_cells("A3:Q3")
    ws["A1"].font=Font(bold=True,color="FFFFFF",size=16); ws["A1"].fill=PatternFill("solid",fgColor="17365D"); ws["A1"].alignment=Alignment(horizontal="center",vertical="center")
    ws["A2"].font=Font(bold=True,color="1F4E78",size=11); ws["A2"].alignment=Alignment(horizontal="center",vertical="center")
    ws["A3"].font=Font(italic=True,color="595959",size=9); ws["A3"].alignment=Alignment(horizontal="center",vertical="center")
    ws.sheet_view.showGridLines=False; ws.freeze_panes="F6"; ws.auto_filter.ref=f"A5:Q5"; ws.row_dimensions[1].height=28; ws.row_dimensions[5].height=30
    for c in ws[5]:
        c.font=Font(bold=True,color="FFFFFF",size=10); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=BORDER

    # Business grouping: OP Type -> Engine Type -> Vehicle No.
    ordered=sorted(roster,key=lambda v:(str(roster[v].get("op","")).upper(),str(roster[v].get("engine","")).upper(),v))
    row=6; previous_group=None
    fy_fills={"2024-25":"F7F9FC","2025-26":"EDF3F8","2026-27":"E2F0D9"}
    for idx,v in enumerate(ordered,1):
        group=(str(roster[v].get("op","")).strip(),str(roster[v].get("engine","")).strip())
        start=row
        # Visually identify each OP Type -> Engine Type group without adding
        # synthetic business data to the report.
        if previous_group is not None and group != previous_group:
            for col in range(1,18):
                ws.cell(row,col).border=Border(top=Side(style="medium",color="4472C4"))
        for fy in FYS:
            vals=[values.get(v,{}).get(fy,{}).get(m,"") for m in MONTHS]; ws.append([idx,v,group[0],group[1],fy,*vals])
            for c in ws[row]:
                c.border=BORDER; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            for c in range(5,18):ws.cell(row,c).fill=PatternFill("solid",fgColor=fy_fills[fy])
            ws.cell(row,5).font=Font(bold=(fy=="2026-27"),color="274E13" if fy=="2026-27" else "000000")
            for c in range(6,18):
                val=ws.cell(row,c).value
                if isinstance(val,(int,float)):ws.cell(row,c).number_format="0.00"
                if isinstance(val,str) and "🔧" in val:ws.cell(row,c).fill=PatternFill("solid",fgColor="FFF2CC");ws.cell(row,c).font=Font(bold=True,color="7F6000")
                elif isinstance(val,str) and "⚙" in val:ws.cell(row,c).fill=PatternFill("solid",fgColor="D9EAF7");ws.cell(row,c).font=Font(bold=True,color="1F4E78")
            row+=1
        end=row-1
        for col in range(1,5):
            ws.merge_cells(start_row=start,start_column=col,end_row=end,end_column=col);ws.cell(start,col).alignment=Alignment(horizontal="center" if col!=3 and col!=4 else "left",vertical="center",wrap_text=True)
        # Strong boundary between vehicle blocks; stronger still when OP/engine group changes.
        next_group=None
        if idx<len(ordered):
            nv=ordered[idx];next_group=(str(roster[nv].get("op","")).strip(),str(roster[nv].get("engine","")).strip())
        style="medium" if next_group!=group else "thin"; bottom=Side(style=style,color="5B9BD5")
        for col in range(1,18):
            cell=ws.cell(end,col);cell.border=Border(left=cell.border.left,right=cell.border.right,top=cell.border.top,bottom=bottom)
        previous_group=group

    # Keep the identifiers visually anchored while the month history scrolls.
    ws.column_dimensions["B"].width=15
    ws.column_dimensions["C"].width=18
    ws.column_dimensions["D"].width=18
    ws.column_dimensions["E"].width=11
    for r in range(6,row):
        ws.row_dimensions[r].height=22
        # Current FY is the decision row: emphasize it without changing values.
        if ws.cell(r,5).value=="2026-27":
            for c in range(1,18):
                ws.cell(r,c).font=Font(
                    name=ws.cell(r,c).font.name or "Calibri",
                    size=ws.cell(r,c).font.sz or 11,
                    bold=True,
                    color=ws.cell(r,c).font.color if ws.cell(r,c).font.color and ws.cell(r,c).font.color.type=="rgb" else "274E13"
                )

    widths=[7,14,20,18,10]+[12]*12
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    row+=1;ws.cell(row,1,"Legend / Data Remarks").font=Font(bold=True,color="FFFFFF");ws.cell(row,1).fill=PatternFill("solid",fgColor="1F4E78");ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=17)
    row+=1;ws.cell(row,1,"🔧 = Schedule-III completed | ⚙ = Schedule-IV completed | number = completion day");ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=17)
    if remarks:
        row+=2;ws.cell(row,1,"DATA EXCEPTIONS").font=Font(bold=True,color="9C0006");ws.cell(row,1).fill=PatternFill("solid",fgColor="F4CCCC");ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=17)
        row+=1;ws.append(["Month","Vehicle","Source maintenance entries - review required"])
        for c in ws[row]:c.font=Font(bold=True);c.fill=PatternFill("solid",fgColor="D9EAF7")
        for month,vehicle,detail in remarks:ws.append([month,vehicle,detail])
    return ws
