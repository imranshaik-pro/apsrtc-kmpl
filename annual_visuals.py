"""Shared Annual KPI presentation model for the existing XLSX/Sheets writers.

No source calculations live here. Blank and MANUAL values stay unavailable.
Completed-year Upto and selected-year YTD are labelled as different periods.
"""
from datetime import datetime

NAVY = "12345B"
BLUE = "1976B9"
GREEN = "287B59"
ORANGE = "B66521"
PURPLE = "7056A5"
INK = "23364D"
PALE = "F3F6FA"
LINE = "CCD7E3"
FY_COLORS = ("EAF2FA", "EAF4EE", "FFF2D9")


def number(value):
    if value is None or value == "" or isinstance(value, bool): return None
    if isinstance(value, (int, float)): return value
    try: return float(str(value).replace(",", ""))
    except ValueError: return None


def records(mat, fys):
    """Read labels even when Sheets omits trailing blank cells."""
    result = {}; name = ""
    for raw in mat:
        row = list(raw) + [""] * max(0, 18-len(raw))
        if str(row[2]) not in fys: continue
        if str(row[1] or "").strip(): name = str(row[1]).strip()
        if name: result.setdefault(name, {})[str(row[2])] = row
    return result


def group(name):
    if name.startswith("HSD"): return "Fuel efficiency (KMPL)", BLUE
    if name.startswith("PRODUCT:"): return "Product KMPL", GREEN
    if name.startswith("ENGINE:"): return "Engine KMPL", ORANGE
    if name in {"AVG TYRE LIFE", "NEW TYRE LIFE", "RC TYRE LIFE", "N.T.S RATE", "Ist RC S Rate", "TTL SCP Rate", "RT Factor"}:
        return "Tyre performance", PURPLE
    return "Other operating KPIs", NAVY


def period_label(selected):
    return datetime.strptime(selected, "%Y-%m").strftime("%B %Y")


def dashboard_model(display, fys, mat, selected):
    data = records(mat, fys)
    model = {"cells": [], "widths": [67]*4+[100]*3+[22]+[63]*10,
             "charts": [], "height": 23, "last_row": 48}
    def cell(r,c,text,rs=1,cs=1,fill="FFFFFF",color=INK,size=10,bold=False,align="left"):
        model["cells"].append(dict(r=r,c=c,text=text,rs=rs,cs=cs,fill=fill,color=color,size=size,bold=bold,align=align))
    period = period_label(selected)
    current = fys[-1]
    endmon = datetime.strptime(selected, "%Y-%m").strftime("%b")
    cell(0,0,"APSRTC  ANNUAL KPI",cs=18,fill=NAVY,color="FFFFFF",size=20,bold=True)
    cell(1,0,f"{display} DEPOT    Reporting through {period}",cs=18,size=14,bold=True)
    cell(2,0,f"FY {fys[0]} and {fys[1]}: full-year Upto     FY {current}: April–{endmon} YTD",cs=18,size=10)
    cards=[("HSD KMPL EXCL AC","HSD EXCL. AC · YTD KMPL",BLUE),
           ("AVG TYRE LIFE","AVERAGE TYRE LIFE · YTD LAKH KM",PURPLE),
           ("B.D RATE","BREAKDOWN RATE · YTD",ORANGE)]
    for i,(key,label,color) in enumerate(cards):
        col=i*6
        val=number(data.get(key,{}).get(current,[""]*18)[17])
        cell(4,col,label,cs=5,fill=color,color="FFFFFF",bold=True,size=10)
        cell(5,col,"Not available" if val is None else val,rs=2,cs=5,fill=PALE,color=color,size=24,bold=True,align="center")
        cell(7,col,f"FY {current} through {endmon}",cs=5,fill=PALE,size=10,align="center")
    cell(9,0,"KPI summary",cs=7,size=13,bold=True)
    cell(10,0,"KPI / parameter",cs=4,fill=NAVY,color="FFFFFF",bold=True)
    for i,fy in enumerate(fys):
        cell(10,4+i,fy+(" YTD" if i==2 else " Full FY"),fill=NAVY,color="FFFFFF",bold=True,align="center")
    previous=None; r=11
    for name, years in data.items():
        title,color=group(name)
        if title!=previous:
            cell(r,0,title,cs=7,fill=color,color="FFFFFF",bold=True); r+=1; previous=title
        label=name.replace("PRODUCT: ","").replace("ENGINE: ","")
        if name in {"AVG TYRE LIFE","NEW TYRE LIFE","RC TYRE LIFE"}: label += " (lakh km)"
        cell(r,0,label,cs=4,fill=PALE if r%2==0 else "FFFFFF")
        for j,fy in enumerate(fys):
            val=number(years.get(fy,[""]*18)[17])
            cell(r,4+j,val if val is not None else "—",fill=FY_COLORS[j],bold=(j==2),align="right")
        r+=1
    # Same-unit time series; no fleet availability or depot-vs-APSRTC comparison.
    sy,sm=map(int,selected.split("-")); count=(sm-4)%12+1
    months=["Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec","Jan","Feb","Mar"][:count]
    for idx,(title,keys,labels,unit) in enumerate([
        ("Monthly HSD KMPL",["HSD KMPL INCL AC","HSD KMPL EXCL AC"],["Incl. AC","Excl. AC"],"KMPL"),
        ("Monthly tyre life",["AVG TYRE LIFE","NEW TYRE LIFE","RC TYRE LIFE"],["Average","New","RC"],"lakh km")]):
        hr=idx*16
        cell(hr,19,"Month")
        for j,label in enumerate(labels,20): cell(hr,j,label)
        for mi,mon in enumerate(months):
            cell(hr+mi+1,19,mon)
            for j,key in enumerate(keys,20):
                source=data.get(key,{}).get(current,[""]*18)
                cell(hr+mi+1,j,number(source[4+mi]))
        model["charts"].append(dict(title=f"{title} — FY {current}",unit=unit,
            source_row=hr,points=count,series=len(keys),row=10+idx*16,col=8,width=630,height=340))
    note=max(r+1,44)
    cell(note,0,"— = missing source value; not zero. Full-year and YTD totals are not like-for-like comparisons.",cs=18,size=10)
    cell(note+1,0,"Tyre life is shown in lakh km. Rate values retain APSRTC source units. Charts stop at the selected month.",cs=18,size=10)
    model["last_row"]=note+2
    return model


def render_xlsx_dashboard(wb,title,model,display,selected):
    # Integrates with the repository's established openpyxl report writer.
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter
    if title in wb.sheetnames: del wb[title]
    ws=wb.create_sheet(title,0); ws.sheet_view.showGridLines=False
    for i,width in enumerate(model["widths"],1): ws.column_dimensions[get_column_letter(i)].width=(width-5)/7
    for r in range(1,model["last_row"]+1): ws.row_dimensions[r].height=18
    ws.row_dimensions[1].height=34; ws.row_dimensions[2].height=26
    for spec in model["cells"]:
        r,c=spec["r"]+1,spec["c"]+1; er,ec=r+spec["rs"]-1,c+spec["cs"]-1
        if er>r or ec>c: ws.merge_cells(start_row=r,start_column=c,end_row=er,end_column=ec)
        for row in ws.iter_rows(min_row=r,max_row=er,min_col=c,max_col=ec):
            for v in row: v.fill=PatternFill("solid",fgColor=spec["fill"])
        v=ws.cell(r,c,spec["text"])
        v.font=Font(name="Arial",size=spec["size"],bold=spec["bold"],color=spec["color"])
        v.alignment=Alignment(horizontal=spec["align"],vertical="center",wrap_text=True)
        v.number_format="0.00"
        if 11<=spec["r"]<model["last_row"]-2 and spec["c"]<7:
            v.border=Border(bottom=Side(style="hair",color=LINE))
    for spec in model["charts"]:
        chart=LineChart(); chart.title=spec["title"]; chart.y_axis.title=spec["unit"]
        chart.style=13; chart.height=9; chart.width=16.6; chart.display_blanks="gap"
        chart.smooth=False
        chart.y_axis.numFmt="0.00"
        chart.visible_cells_only=False
        start=spec["source_row"]+1
        chart.add_data(Reference(ws,min_col=21,max_col=20+spec["series"],min_row=start,max_row=start+spec["points"]),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=20,min_row=start+1,max_row=start+spec["points"]))
        chart.legend.position="b"
        for ser,color in zip(chart.series,[BLUE,GREEN,ORANGE]):
            ser.smooth=False
            ser.graphicalProperties.line.solidFill=color
            ser.graphicalProperties.line.width=24000
        ws.add_chart(chart,f"I{spec['row']+1}")
    for col in range(20,24): ws.column_dimensions[get_column_letter(col)].hidden=True
    print_setup(ws,display,period_label(selected),"1:3",18,model["last_row"],one_page=True)
    return ws


def print_setup(ws,display,period,repeat,last_col,last_row,one_page=False):
    from openpyxl.utils import get_column_letter
    ws.print_title_rows=repeat
    ws.print_area=f"A1:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.oddHeader.center.text=f"APSRTC — {display} DEPOT — {period}"
    ws.oddFooter.left.text=f"{display} DEPOT — {period}"
    ws.oddFooter.right.text="Page &P of &N"


def rgb(hexcode):
    return dict(zip(("red","green","blue"),(int(hexcode[i:i+2],16)/255 for i in (0,2,4))))


def google_dashboard_requests(sid,model):
    def region(r,c,rs=1,cs=1):
        return {"sheetId":sid,"startRowIndex":r,"endRowIndex":r+rs,"startColumnIndex":c,"endColumnIndex":c+cs}
    req=[{"updateSheetProperties":{"properties":{"sheetId":sid,"gridProperties":{"rowCount":max(100,model['last_row']),"columnCount":24,"frozenRowCount":0,"frozenColumnCount":0,"hideGridlines":True}},"fields":"gridProperties"}},
         {"unmergeCells":{"range":region(0,0,100,24)}},
         {"updateCells":{"range":region(0,0,100,24),"fields":"userEnteredValue,userEnteredFormat"}}]
    for i,width in enumerate(model["widths"]):
        req.append({"updateDimensionProperties":{"range":{"sheetId":sid,"dimension":"COLUMNS","startIndex":i,"endIndex":i+1},"properties":{"pixelSize":width},"fields":"pixelSize"}})
    req.append({"updateDimensionProperties":{"range":{"sheetId":sid,"dimension":"ROWS","startIndex":0,"endIndex":model["last_row"]},"properties":{"pixelSize":model["height"]},"fields":"pixelSize"}})
    for r,height in [(0,45),(1,34)]:
        req.append({"updateDimensionProperties":{"range":{"sheetId":sid,"dimension":"ROWS","startIndex":r,"endIndex":r+1},"properties":{"pixelSize":height},"fields":"pixelSize"}})
    for s in model["cells"]:
        area=region(s['r'],s['c'],s['rs'],s['cs'])
        if s['rs']>1 or s['cs']>1: req.append({"mergeCells":{"range":area,"mergeType":"MERGE_ALL"}})
        fmt={"backgroundColor":rgb(s['fill']),"textFormat":{"fontFamily":"Arial","fontSize":s['size'],"bold":s['bold'],"foregroundColor":rgb(s['color'])},"verticalAlignment":"MIDDLE","horizontalAlignment":s['align'].upper(),"wrapStrategy":"WRAP","numberFormat":{"type":"NUMBER","pattern":"0.00"}}
        req.append({"repeatCell":{"range":area,"cell":{"userEnteredFormat":fmt},"fields":"userEnteredFormat"}})
        v=s['text']; value={} if v is None else {"numberValue":v} if isinstance(v,(float,int)) else {"stringValue":v}
        req.append({"updateCells":{"start":{"sheetId":sid,"rowIndex":s['r'],"columnIndex":s['c']},"rows":[{"values":[{"userEnteredValue":value}]}],"fields":"userEnteredValue"}})
    for s in model["charts"]:
        r=s['source_row']; count=s['points']+1
        source=lambda c:{"sourceRange":{"sources":[region(r,c,count)]}}
        basic={"chartType":"LINE","headerCount":1,"legendPosition":"BOTTOM_LEGEND",
               "domains":[{"domain":source(19)}],"series":[{"series":source(20+j),"targetAxis":"LEFT_AXIS","color":rgb([BLUE,GREEN,ORANGE][j])} for j in range(s['series'])],
               "axis":[{"position":"LEFT_AXIS","title":s['unit']}],"lineSmoothing":False}
        chart={"spec":{"title":s['title'],"fontName":"Arial","hiddenDimensionStrategy":"SHOW_ALL","basicChart":basic},
               "position":{"overlayPosition":{"anchorCell":{"sheetId":sid,"rowIndex":s['row'],"columnIndex":s['col']},"widthPixels":s['width'],"heightPixels":s['height']}}}
        req.append({"addChart":{"chart":chart}})
    req.append({"updateDimensionProperties":{"range":{"sheetId":sid,"dimension":"COLUMNS","startIndex":19,"endIndex":24},"properties":{"hiddenByUser":True},"fields":"hiddenByUser"}})
    req.append({"updateSheetProperties":{"properties":{"sheetId":sid,"gridProperties":{"frozenRowCount":3}},"fields":"gridProperties.frozenRowCount"}})
    return req
