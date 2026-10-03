"""Read-only annual workbook verification. No network, refresh or source repair.

Different report endpoints are compared as diagnostics, never substituted.
Checks against a saved snapshot establish preservation, not live correctness.
"""
import argparse
import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import annual_history as history
from annual_visuals import number, records


def count_sum(*values):
    return sum(values) if all(isinstance(v,(int,float)) and not isinstance(v,bool) for v in values) else None


def ratio_reviews(mat, fys, selected):
    """Screen complete monthly coverage; absence of weights prevents a proof."""
    result = []
    count = (int(selected[-2:]) - 4) % 12 + 1
    current = history.fy_of(selected)
    for name, years in records(mat, fys).items():
        if not name.startswith(('HSD', 'PRODUCT:', 'ENGINE:')):
            continue
        row = years.get(current, [''] * 18)
        values = [number(v) for v in row[4:4+count]]
        upto = number(row[17])
        if count < 2 or upto is None or any(v is None for v in values):
            result.append(dict(check='Monthly range', kpi=name, severity='unverified',
                               reason='Missing cumulative value or complete monthly coverage'))
        elif Decimal(str(upto)) < Decimal(str(min(values)))-Decimal('.01') or Decimal(str(upto)) > Decimal(str(max(values)))+Decimal('.01'):
            result.append(dict(check='Monthly range', kpi=name, severity='review',
                               observed=upto, monthly_min=min(values), monthly_max=max(values),
                               reason='Outside rounded monthly range; verify identical scope and positive weights in original source'))
    return result


def audit_workbook(path, depot, selected, calculated_path=None):
    from openpyxl import load_workbook
    import report_details as details
    wb = load_workbook(path)
    calculated = load_workbook(calculated_path or path, data_only=True)
    fys = [f'{y}-{str(y+1)[-2:]}' for y in range(int(history.fy_of(selected)[:4])-2,
                                               int(history.fy_of(selected)[:4])+1)]
    mat = [list(r) for r in wb['2. Detailed Data (Our Format)'].values]
    data = records(mat, fys)
    report = dict(depot=depot, selected=selected, file_sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                  checks=[], reviews=ratio_reviews(mat, fys, selected), limits=[])
    def check(kind, name, actual, expected, cell):
        passed = actual == expected
        if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
            passed = abs(actual-expected) < 1e-10
        report['checks'].append(dict(check=kind, kpi=name, cell=cell, actual=actual,
                                     expected=expected, result='pass' if passed else 'fail'))
    dashboard = wb['1. KPI Dashboard']
    year_cells = {}
    for row in dashboard:
        for cell in row:
            for fy in fys:
                if str(cell.value).startswith(fy+' '):
                    year_cells[fy] = cell.column
    for name, years in data.items():
        label = name.replace('PRODUCT: ', '').replace('ENGINE: ', '')
        if name in {'AVG TYRE LIFE','NEW TYRE LIFE','RC TYRE LIFE'}:
            label += ' (lakh km)'
        match = next((c for row in dashboard for c in row if c.value == label), None)
        if match is None:
            # The compact dashboard omits unavailable engines; source detail
            # rows remain present and are compared separately.
            continue
        for fy in fys:
            if fy not in year_cells:
                continue
            cell = dashboard.cell(match.row, year_cells[fy])
            check('Dashboard to original detail', name, number(cell.value),
                  number(years.get(fy, ['']*18)[17]), f'{dashboard.title}!{cell.coordinate}')
    heroes = {'HSD INCL. AC · YTD KMPL':'HSD KMPL INCL AC',
              'HSD EXCL. AC · YTD KMPL':'HSD KMPL EXCL AC',
              'AVG TYRE LIFE · LAKH KM':'AVG TYRE LIFE',
              'AVERAGE TYRE LIFE · YTD LAKH KM':'AVG TYRE LIFE',
              'BREAKDOWN RATE · YTD':'B.D RATE'}
    for row in dashboard:
        for cell in row:
            if cell.value in heroes:
                name = heroes[cell.value]
                check('Hero to original detail', name, number(dashboard.cell(cell.row+1,cell.column).value),
                      number(data.get(name,{}).get(fys[-1],['']*18)[17]),
                      f'{dashboard.title}!{dashboard.cell(cell.row+1,cell.column).coordinate}')
    if history.CACHE_TITLE not in wb.sheetnames:
        report['limits'].append('Original KPI _ANNUAL_HISTORY is absent: no original HSD/Product/Engine/LUB/BD/Tyre-life source provenance or weights can be certified.')
    else:
        history.decode([list(r) for r in wb[history.CACHE_TITLE].values], depot)
    cache = details.decode([list(r) for r in wb[details.CACHE_TITLE].values], depot)
    report['detail_cache_checksum'] = 'verified'
    for period in details.annual_periods(selected):
        groups = cache['months'].get(period,{})
        metrics = {g:dict(zip(s['headers'],s['rows'][0])) for g,s in groups.items() if g in ('B','F')}
        b,f = metrics.get('B',{}),metrics.get('F',{})
        for label,observed,numerator,denominator in (
            ('Mech. Defects %',b.get('Mech. Defetcs %'),b.get('Total Mech Defects'),b.get('Num. of Tyres received')),
            ('Stone %',f.get('Stone %'),f.get('Stones'),f.get('Number of Tyres Received')),
            ('Worn Smooth %',f.get('Worn Smooth %'),
             count_sum(f.get('Worn smooth RC'),f.get('Worn smooth RT')),
             count_sum(f.get('Total RC'),f.get('Total RT')))):
            if any(not isinstance(v,(int,float)) for v in (observed,numerator,denominator)) or denominator<=0:
                continue
            expected = float((Decimal(str(numerator))*100/Decimal(str(denominator))).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
            if abs(observed-expected)>1e-10:
                report['reviews'].append(dict(check='Source percentage arithmetic',kpi=label,period=period,
                    observed=observed,comparison=expected,severity='review',reason='Reported percentage differs from captured counts; source value retained'))
            else:
                check('Source percentage arithmetic',label,observed,expected,period)
    for name, years in data.items():
        row = years.get(fys[-1], ['']*18)
        count = (int(selected[-2:])-4) % 12+1
        for i in range(count,12):
            actual = row[4+i]
            check('Future month hidden from view', name, None if actual in ('',None) else actual,
                  None, f'Detailed Data {history.MONTHS[i]}')
    fixture_root = Path(__file__).parent/'tests'/'fixtures'/'report-details'/depot
    for group, snapshot in cache['months'].get(selected,{}).items():
        fixture = fixture_root/f'{selected}-{group}.html'
        if not fixture.exists():
            continue
        parsed = (details.parse_engine(fixture.read_text(),depot,selected,group)
                  if group in ('UD','UM') else
                  details.parse_tyre(fixture.read_text(),group,snapshot['tyre_code'],selected,depot))
        for field in ('headers','rows'):
            check('Captured HTML to saved snapshot', f'{group} {field}',snapshot[field],parsed[field],
                  str(fixture.relative_to(Path(__file__).parent)))
    for title, blocks in details.presentation_sections(cache, selected, True).items():
        ws, calc = wb[title], calculated[title]
        for block in blocks:
            anchor = next((r for r in range(1, ws.max_row+1) if ws.cell(r, 1).value == block['title']), None)
            if anchor is None:
                check('Section present', block['title'], None, 'present', title)
                continue
            start = anchor + 3 + bool(block.get('header_groups'))
            for offset, row in enumerate(block['rows']):
                if offset == block.get('total_index'):
                    continue
                for c, expected in enumerate(row, 1):
                    actual = calc.cell(start+offset, c).value
                    if actual is None and expected in ('', None):
                        actual = expected
                    check('Saved source to detail view', block['title'], actual, expected,
                          f'{title}!{ws.cell(start+offset,c).coordinate}')
            if 'total_index' in block:
                n = block['total_index']; tr = start+n
                for c in range(3, len(block['headers'])+1):
                    if c in block.get('percent_columns', []):
                        check('Percentage FY total blank', block['title'], ws.cell(tr,c).value, None,
                              f'{title}!{ws.cell(tr,c).coordinate}')
                        continue
                    letter = ws.cell(tr,c).column_letter
                    expected = f'=IF(COUNT({letter}{start}:{letter}{tr-1})={n},SUM({letter}{start}:{letter}{tr-1}),"")'
                    check('Count FY total formula', block['title'], ws.cell(tr,c).value, expected,
                          f'{title}!{ws.cell(tr,c).coordinate}')
                    vals = [r[c-1] for r in block['rows'][:n]]
                    total = sum(vals) if all(isinstance(v,(int,float)) for v in vals) else None
                    check('Count FY total calculation', block['title'], calc.cell(tr,c).value, total,
                          f'{title}!{ws.cell(tr,c).coordinate}')
    # Matching exact product names only. Separate endpoints have no established
    # equivalence contract, so a difference is a review, never an automatic fix.
    snap = cache['months'].get(selected, {}).get('UM')
    if snap:
        total = next((row for row in snap['rows'] if row[0] == 'TOTAL'), None)
        if total:
            for label, value in zip(snap['headers'][2:], total[2:]):
                key = 'PRODUCT: '+label
                if key not in data:
                    continue
                actual = number(data[key].get(fys[-1], ['']*18)[17])
                if actual != number(value):
                    report['reviews'].append(dict(check='Separate endpoint comparison', kpi=key,
                        severity='review', observed=actual, comparison=value,
                        reason='prodkmpl_um.php vs mednew/eng_prod_um.php UM; source equivalence and original response not available'))
    for period in details.annual_periods(selected):
        snap = cache['months'].get(period, {}).get('C')
        if not snap:
            continue
        values = dict(zip(snap['headers'], snap['rows'][0]))
        for stage in details.SCRAP_GROUPS:
            codes = [values[f'{stage} — S{i}'] for i in range(1,10)]
            total = values[f'{stage} — Total']
            if all(isinstance(v,(int,float)) for v in codes) and sum(codes) != total:
                report['reviews'].append(dict(check='Statement C source arithmetic', kpi=stage,
                    period=period, severity='review', observed=total, stage_sum=sum(codes),
                    reason='Source inconsistency retained; do not replace reported total'))
    report['summary'] = dict(passed=sum(c['result']=='pass' for c in report['checks']),
                            failed=sum(c['result']=='fail' for c in report['checks']),
                            review=sum(r['severity']=='review' for r in report['reviews']),
                            unverified=sum(r['severity']=='unverified' for r in report['reviews']))
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook'); parser.add_argument('--depot', required=True)
    parser.add_argument('--month', required=True); parser.add_argument('--calculated-workbook')
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    report = audit_workbook(args.workbook, args.depot, args.month, args.calculated_workbook)
    Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(report['summary']))
    return bool(report['summary']['failed'])


if __name__ == '__main__':
    raise SystemExit(main())
