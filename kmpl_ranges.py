"""Verified APSRTC monthly range counts, with separate durable FY history.

These are source classifications, not buckets recalculated from daily KMPL.
The existing workbook/Sheets serializers are deliberately reused.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
from datetime import date, datetime

from bs4 import BeautifulSoup
import report_details as d

CACHE_TITLE = "_KMPL_RANGE_HISTORY"
SCHEMA = "kmpl-range-2"
MONTHLY_TITLE = "KMPL Range Distribution"
ANNUAL_TITLE = "FY KMPL Range Trend"
ENDPOINTS = {"VEHICLE": "vehlog_kmpldepot.php", "DRIVER": "drvlog_kmpldepot.php"}
UPTO_SELECTORS = {"VEHICLE": "vehlog_kmpl.php", "DRIVER": "drvlog_ckmpl.php"}
UPTO_ENDPOINTS = {"VEHICLE": "vehlog_kmpldepot.php", "DRIVER": "drvlog_ckmpldepot.php"}
SCOPES = ("MONTH", "UPTO")
LABELS = ("Below -3.00", "3.00 - 4.00", "4.01 - 5.00", "5.01 - 5.15",
          "5.16 - 5.30", "5.31 - 5.60", "5.60 - ABV")
HEADER_KEYS = tuple(d.key(x) for x in LABELS) + ("GTOTAL",)


def count(value):
    text = str(value).strip().replace(",", "")
    if not re.fullmatch(r"\d+", text):
        raise ValueError(f"Range count must be a non-negative integer: {value!r}")
    return int(text)


def parse_range(html, entity, depot, period, region, scope="MONTH"):
    if entity not in ENDPOINTS:
        raise ValueError("Unknown range source")
    soup = BeautifulSoup(html, "html.parser")
    headings = [d.norm(h.get_text(" ", strip=True)) for h in soup.find_all(re.compile(r"^h[1-6]$"))]
    expected = d.month_end(period).strftime("%m/%Y")
    if not any(re.search(rf"RANGE WISE {entity} HSD KMPL FOR THE MONTH OF:\s*{expected}\s*$", h)
               for h in headings):
        raise ValueError("Range source entity/month heading is not verified")
    choices = []
    for table in soup.find_all("table"):
        trs = table.find_all("tr")
        if len(trs) < 3:
            continue
        first = trs[0].find_all(["th", "td"], recursive=False)
        second = trs[1].find_all(["th", "td"], recursive=False)
        if len(first) != 3 or len(second) != 8:
            continue
        if (d.key(first[0].get_text()) != "SLNO" or d.key(first[1].get_text()) not in ("REGION", "DEPOT")
                or first[0].get("rowspan") != "2" or first[1].get("rowspan") != "2"
                or first[2].get("colspan") != "8"
                or d.key(first[2].get_text()) != f"KMPLRANGEWISENOOF{entity}S"
                or tuple(d.key(c.get_text(" ", strip=True)) for c in second) != HEADER_KEYS):
            continue
        rows, seen, total = [], set(), None
        for tr in trs[2:]:
            cells = tr.find_all(["td", "th"], recursive=False)
            if len(cells) != 10:
                raise ValueError("Range row width does not match source headings")
            values = [c.get_text(" ", strip=True) for c in cells]
            counts = [count(v) for v in values[2:]]
            if sum(counts[:7]) != counts[7]:
                raise ValueError("Range row does not reconcile to G.Total")
            name = d.norm(values[1])
            if name == "TOTAL":
                if total is not None or tr is not trs[-1]:
                    raise ValueError("Expected one final regional total")
                total = counts
            else:
                if not re.fullmatch(r"\d+", values[0]) or not name or name in seen or total is not None:
                    raise ValueError("Duplicate/invalid depot range identity")
                link = cells[1].find("a", href=True)
                if link:
                    from urllib.parse import parse_qs, urlparse
                    query = parse_qs(urlparse(link["href"]).query)
                    if query.get("yymm") != [period.replace("-", "")] or query.get("rreg") != [name]:
                        raise ValueError("Range drilldown depot/month does not match its row")
                seen.add(name)
                rows.append([int(values[0]), name, *counts])
        if not rows or total is None or any(sum(row[i+2] for row in rows) != total[i] for i in range(8)):
            raise ValueError("Regional range totals do not reconcile")
        selected = [row for row in rows if row[1] == d.norm(depot)]
        if len(selected) != 1:
            raise ValueError("No unique exact selected-depot range row")
        choices.append(dict(depot=d.norm(depot), period=period, entity=entity, scope=scope, region=region,
                            headers=["SL No", "Depot", *[c.get_text(" ", strip=True) for c in second]],
                            rows=rows, regional_total=total, counts=selected[0][2:],
                            source=f"{d.BASE}/med/{(ENDPOINTS if scope == 'MONTH' else UPTO_ENDPOINTS)[entity]}"))
    if len(choices) != 1:
        layouts=[]
        for table in soup.find_all("table"):
            layouts.append([[dict(text=c.get_text(" ",strip=True),key=d.key(c.get_text()),rowspan=c.get("rowspan"),colspan=c.get("colspan"))
                             for c in tr.find_all(["th","td"],recursive=False)] for tr in table.find_all("tr")[:2]])
        raise ValueError("Expected one verified range table; empty/login pages are unavailable; expected="+str((entity,HEADER_KEYS))+"; headers="+json.dumps(layouts))
    return choices[0]


class SourceAdapter:
    def __init__(self, session, depot, region, today=None):
        self.session, self.depot, self.region = session, depot, region
        self.today = today or datetime.now(d.IST).date()

    def __call__(self, entity, period, scope="MONTH"):
        if d.month_end(period).replace(day=1) > self.today:
            raise ValueError("Future range months are not requested")
        payload = dict(action="", yymm=period.replace("-", ""), rreg=self.region)
        if scope not in SCOPES: raise ValueError("Unknown range scope")
        if scope == "UPTO":
            selector = self.session.get(f"{d.BASE}/med/{UPTO_SELECTORS[entity]}", timeout=45)
            selector.raise_for_status()
        endpoints = ENDPOINTS if scope == "MONTH" else UPTO_ENDPOINTS
        response = self.session.get(f"{d.BASE}/med/{endpoints[entity]}", params=payload, timeout=45)
        response.raise_for_status()
        snap = parse_range(response.text, entity, self.depot, period, self.region, scope)
        snap.update(request=payload, fetched_at=datetime.now(d.IST).isoformat(timespec="seconds"),
                    provisional=d.month_end(period) >= self.today)
        return snap


def new_cache(depot):
    return dict(schema=SCHEMA, depot=d.norm(depot), months={}, errors={})


def validate_snapshot(snap, depot, entity, period, scope="MONTH"):
    if (not isinstance(snap, dict) or snap.get("depot") != d.norm(depot)
            or snap.get("entity") != entity or snap.get("period") != period or snap.get("scope", "MONTH") != scope):
        raise ValueError("Range snapshot identity mismatch")
    counts = snap.get("counts", [])
    if len(counts) != 8 or any(type(n) is not int or n < 0 for n in counts) or sum(counts[:7]) != counts[7]:
        raise ValueError("Invalid cached range counts")
    rows = snap.get("rows", [])
    names = [r[1] for r in rows if isinstance(r, list) and len(r) == 10]
    if len(names) != len(rows) or len(set(names)) != len(names):
        raise ValueError("Invalid cached range depot rows")
    for row in rows:
        if (type(row[0]) is not int or not isinstance(row[1], str)
                or any(type(n) is not int or n < 0 for n in row[2:]) or sum(row[2:9]) != row[9]):
            raise ValueError("Invalid cached range row totals")
    chosen = [row[2:] for row in rows if row[1] == d.norm(depot)]
    total = snap.get("regional_total", [])
    if chosen != [counts] or total != [sum(r[i+2] for r in rows) for i in range(8)]:
        raise ValueError("Cached regional/depot range totals mismatch")
    if len(snap.get("headers", [])) != 10:
        raise ValueError("Cached range headings missing")
    return snap


def encode(cache):
    raw = json.dumps(cache, ensure_ascii=True, allow_nan=False, separators=(",", ":"))
    chunks = [raw[i:i+30000] for i in range(0, len(raw), 30000)]
    return [[SCHEMA, cache["depot"], len(chunks), hashlib.sha256(raw.encode()).hexdigest()]] + [[v] for v in chunks]


def decode(rows, depot):
    if not rows or len(rows[0]) < 4 or list(rows[0][:2]) != [SCHEMA, d.norm(depot)]:
        raise ValueError("Range history identity/schema mismatch")
    n = int(rows[0][2])
    if n < 1 or len(rows) < n+1 or any(not r or not isinstance(r[0], str) for r in rows[1:n+1]):
        raise ValueError("Range history chunks incomplete")
    raw = "".join(r[0] for r in rows[1:n+1])
    if hashlib.sha256(raw.encode()).hexdigest() != rows[0][3]:
        raise ValueError("Range history checksum mismatch")
    cache = json.loads(raw)
    if cache.get("schema") != SCHEMA or cache.get("depot") != d.norm(depot):
        raise ValueError("Range history identity mismatch")
    for period, groups in cache["months"].items():
        d.month_end(period)
        for key, snap in groups.items():
            parts = key.rsplit("_", 1)
            entity, scope = (parts[0], parts[1]) if len(parts) == 2 and parts[1] in SCOPES else (key, "MONTH")
            if entity not in ENDPOINTS:
                raise ValueError("Unknown cached range entity")
            validate_snapshot(snap, depot, entity, period, scope)
    cache.setdefault("errors", {})
    return cache


def periods_for_fys(fys, selected):
    d.month_end(selected)
    return [f"{year:04d}-{month:02d}" for fy in fys for year, month in
            [(int(fy[:4])+(month < 4), month) for month in list(range(4,13))+[1,2,3]]
            if f"{year:04d}-{month:02d}" <= selected]


def update(cache, periods, fetch, selected, checkpoint=lambda: None, include_upto=False):
    calls = []
    for period in periods:
        d.month_end(period)
        groups = cache["months"].setdefault(period, {})
        scopes = SCOPES if include_upto and period == selected else ("MONTH",)
        for entity in ENDPOINTS:
            for scope in scopes:
                key = f"{entity}_{scope}"
                old = groups.get(key) or (groups.get(entity) if scope == "MONTH" else None)
                if old and not old.get("provisional") and period != selected:
                    groups[key] = old
                    continue
                calls.append((entity, period, scope))
                try:
                    snap = validate_snapshot(fetch(entity, period) if scope == "MONTH" else fetch(entity, period, "UPTO"),
                                             cache["depot"], entity, period, scope)
                    groups[key] = copy.deepcopy(snap)
                    cache["errors"].get(period, {}).pop(key, None)
                except Exception as exc:
                    cache["errors"].setdefault(period, {})[key] = f"{type(exc).__name__}: {exc}"[:300]
                    print(f"RANGE_SOURCE_UNAVAILABLE {cache['depot']} {period} {entity} {scope}: {exc}")
        checkpoint()
    return calls

def read_google_cache(service, sid, depot):
    api = service.spreadsheets()
    meta = api.get(spreadsheetId=sid, fields="sheets.properties").execute()
    if not any(s["properties"]["title"] == CACHE_TITLE for s in meta.get("sheets", [])):
        return new_cache(depot)
    return decode(api.values().get(spreadsheetId=sid, range=f"'{CACHE_TITLE}'!A:D").execute().get("values", []), depot)


def save_google_cache(service, sid, cache):
    api = service.spreadsheets()
    rows = encode(cache)
    meta = api.get(spreadsheetId=sid, fields="sheets.properties").execute()
    existing = next((s for s in meta.get("sheets", []) if s["properties"]["title"] == CACHE_TITLE), None)
    if existing is None:
        api.batchUpdate(spreadsheetId=sid, body={"requests": [{"addSheet": {"properties": {
            "title": CACHE_TITLE, "hidden": True, "gridProperties": {"rowCount": max(100,len(rows)+2), "columnCount": 4}}}}]}).execute()
    else:
        props = existing["properties"]
        api.batchUpdate(spreadsheetId=sid, body={"requests": [{"updateSheetProperties": {"properties": {
            "sheetId": props["sheetId"], "hidden": True, "gridProperties": {
                "rowCount": max(props.get("gridProperties", {}).get("rowCount", 0),len(rows)+2), "columnCount": 4}},
            "fields": "hidden,gridProperties.rowCount,gridProperties.columnCount"}}]}).execute()
    api.values().update(spreadsheetId=sid, range=f"'{CACHE_TITLE}'!A1", valueInputOption="RAW", body={"values": rows}).execute()


def write_xlsx_cache(wb, cache):
    if CACHE_TITLE in wb.sheetnames:
        del wb[CACHE_TITLE]
    ws = wb.create_sheet(CACHE_TITLE)
    for row in encode(cache):
        ws.append(row)
    ws.sheet_state = "hidden"


def render_tab(wb, cache, selected, fys=None):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.series import SeriesLabel
    from report_branding import add_logo
    title = ANNUAL_TITLE if fys else MONTHLY_TITLE
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    for r, text in [(1,f"APSRTC — {cache['depot']} DEPOT"), (2,title.upper())]:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=13)
        ws.cell(r,1,text).font = Font(name="Arial",size=16 if r==1 else 13,bold=True,color="FFFFFF")
        ws.cell(r,1).fill = PatternFill("solid",fgColor=d.NAVY)
        ws.cell(r,1).alignment = Alignment(horizontal="center",vertical="center")
        ws.row_dimensions[r].height = 30
    ws.merge_cells("A3:C3"); ws.merge_cells("D3:M3")
    ws["D3"] = ("FY " + " / ".join(fys) + " · Through " if fys else "Reporting month · ") + d.month_end(selected).strftime("%B %Y")
    ws["D3"].alignment = Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws["D3"].font = Font(name="Arial",size=11,bold=True,color=d.INK)
    ws.merge_cells("A4:M4")
    ws["A4"] = "Source counts; monthly snapshots are not distinct FY populations. Missing ≠ zero. Source 5.60 labels overlap; no reclassification."
    ws["A4"].alignment = Alignment(wrap_text=True,vertical="center")
    ws["A4"].font = Font(name="Arial",size=10,color=d.INK)
    ws.row_dimensions[4].height = 32
    headers = ["Month" if fys else "Population",*LABELS,"G.Total","First 3 ranges count","First 3 ranges %","Highest range %","Source status"]
    r = 6

    def style_row(row, header=False, selected_row=False):
        for cell in ws[row]:
            cell.font = Font(name="Arial",size=10,bold=header or selected_row,color="FFFFFF" if header else d.INK)
            cell.fill = PatternFill("solid",fgColor=d.BLUE if header else "E2F0D9" if selected_row else d.PALE if row%2==0 else "FFFFFF")
            cell.alignment = Alignment(horizontal="left" if cell.column in (1,13) else "right",vertical="center",wrap_text=header or cell.column in (1,13))
            cell.border = Border(bottom=Side(style="thin",color=d.LINE))
        ws.row_dimensions[row].height = 46 if header else 31

    def note(text):
        nonlocal r
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=13)
        ws.cell(r,1,text).font = Font(name="Arial",size=9,italic=True,color="63758A")
        ws.cell(r,1).alignment = Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[r].height = 34
        r += 1

    if not fys:
        note("Selected-depot source counts · Month and Upto/Cum are independent APSRTC source snapshots")
        for c,label in enumerate(["KMPL Range","Vehicle Month","Vehicle Upto","Driver Month","Driver Upto"],1):
            ws.cell(r,c,label)
        style_row(r,header=True); r += 1
        groups=cache["months"].get(selected,{})
        snaps=[groups.get("VEHICLE_MONTH") or groups.get("VEHICLE"), groups.get("VEHICLE_UPTO"),
               groups.get("DRIVER_MONTH") or groups.get("DRIVER"), groups.get("DRIVER_UPTO")]
        for idx,label in enumerate(list(LABELS)+["G.Total"]):
            ws.cell(r,1,label)
            for c,snap in enumerate(snaps,2):
                if snap: ws.cell(r,c,snap["counts"][idx]).number_format="#,##0"
            style_row(r,selected_row=label=="G.Total"); r += 1
        note("Upto/Cum values are source-backed and never calculated by summing monthly populations.")
        r += 1

    for fy in (fys or [None]):
        periods = periods_for_fys([fy],selected) if fy else [selected]
        for entity in ENDPOINTS:
            if r>6: ws.row_breaks.append(Break(id=r-1))
            note(f"{entity.title()} range counts" + (f" · FY {fy}" if fy else "") + f" · {d.BASE}/med/{ENDPOINTS[entity]}")
            for c,label in enumerate(headers,1): ws.cell(r,c,label)
            header_row = r
            style_row(r,header=True); r += 1
            for period in periods:
                snap = (cache["months"].get(period,{}).get(f"{entity}_MONTH") or cache["months"].get(period,{}).get(entity))
                error = cache["errors"].get(period,{}).get(f"{entity}_MONTH") or cache["errors"].get(period,{}).get(entity)
                ws.cell(r,1,d.month_end(period).strftime("%b-%Y") if fy else entity.title())
                if snap:
                    for c,n in enumerate(snap["counts"],2): ws.cell(r,c,n).number_format="#,##0"
                    ws.cell(r,10,f"=SUM(B{r}:D{r})").number_format="#,##0"
                    ws.cell(r,11,f'=IF(I{r}=0,"",J{r}/I{r})').number_format="0.0%"
                    ws.cell(r,12,f'=IF(I{r}=0,"",H{r}/I{r})').number_format="0.0%"
                    status = "Stale — saved source" if error else "Provisional — open month" if snap.get("provisional") else "Verified source"
                    ws.cell(r,13,status)
                else:
                    ws.cell(r,13,"Unavailable — no verified source")
                style_row(r,selected_row=period==selected)
                r += 1
            last_data_row = r-1
            if fys and periods:
                upto_period=periods[-1]
                upto=cache["months"].get(upto_period,{}).get(f"{entity}_UPTO")
                ws.cell(r,1,"Upto/Cum "+d.month_end(upto_period).strftime("%b-%y"))
                if upto:
                    for c,n in enumerate(upto["counts"],2): ws.cell(r,c,n).number_format="#,##0"
                    ws.cell(r,10,f"=SUM(B{r}:D{r})").number_format="#,##0"
                    ws.cell(r,11,f'=IF(I{r}=0,"",J{r}/I{r})').number_format="0.0%"
                    ws.cell(r,12,f'=IF(I{r}=0,"",H{r}/I{r})').number_format="0.0%"
                    ws.cell(r,13,"Provisional — open month" if upto.get("provisional") else "Verified cumulative source")
                else:
                    ws.cell(r,13,"Unavailable — no verified cumulative source")
                style_row(r,selected_row=True); r += 1
            note("Source first three ranges = below 3.00 through the source 4.01–5.00 bucket. Highest range uses source 5.60–ABV. No FY sum of monthly populations.")
            if fys and any((cache["months"].get(p,{}).get(f"{entity}_MONTH") or cache["months"].get(p,{}).get(entity,{})).get("counts",[0]*8)[7]>0 for p in periods):
                chart=LineChart()
                chart.title=f"{entity.title()} · FY {fy} · range shares"
                chart.y_axis.title="Share of monthly source population"
                chart.y_axis.numFmt="0%"
                chart.y_axis.scaling.min=0; chart.y_axis.scaling.max=1
                chart.x_axis.title="Month"; chart.height=7.4; chart.width=28
                chart.add_data(Reference(ws,min_col=11,max_col=12,min_row=header_row,max_row=last_data_row),titles_from_data=True)
                chart.series[0].tx=SeriesLabel(v="First 3 ranges %")
                chart.series[1].tx=SeriesLabel(v="Highest range %")
                chart.set_categories(Reference(ws,min_col=1,min_row=header_row+1,max_row=last_data_row))
                chart.display_blanks="gap"
                chart.series[0].graphicalProperties.line.solidFill="D76A36"
                chart.series[1].graphicalProperties.line.solidFill="287B59"
                for series in chart.series:
                    series.marker.symbol="circle"; series.marker.size=5
                ws.add_chart(chart,f"A{r}")
                for empty_row in range(r,r+14):ws.row_dimensions[empty_row].height=18
                r += 14
            elif fys:
                note("Chart unavailable: no verified non-zero population for this FY. Missing history has not been filled with zeros.")
            if not fys:
                snap = (cache["months"].get(selected,{}).get(f"{entity}_MONTH") or cache["months"].get(selected,{}).get(entity))
                if snap:
                    note(f"Region {snap['region']} · {entity.title()} source table · Fetched {snap.get('fetched_at','not recorded')}")
                    for c,label in enumerate(snap["headers"],1): ws.cell(r,c,label)
                    style_row(r,header=True); r += 1
                    for row in snap["rows"]+[[None,"Regional Total",*snap["regional_total"]]]:
                        for c,n in enumerate(row,1): ws.cell(r,c,n)
                        style_row(r,selected_row=row[1]==cache["depot"])
                        for c in range(3,11): ws.cell(r,c).number_format="#,##0"
                        ws.cell(r,2).alignment = Alignment(wrap_text=True,vertical="center")
                        r += 1
            r += 2
    for c in range(1,14): ws.column_dimensions[get_column_letter(c)].width = 17 if c==1 else 24 if c==13 else 13
    # Region source depot names need more room than a KMPL bucket.
    if not fys: ws.column_dimensions["B"].width = 25
    logo_width = sum(ws.column_dimensions[get_column_letter(c)].width*7+5 for c in (1,2,3))-10
    add_logo(ws,anchor="A3",width=logo_width)
    ws.row_dimensions[3].height = max(40,logo_width*130/640*.75+3)
    ws.freeze_panes = "B5"
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:M{r-2}"
    ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.oddFooter.center.text=f"APSRTC | {cache['depot']} | {selected}"
    ws.oddFooter.right.text="Page &P of &N"
    return ws


def publish_google_tab(service, sid, workbook):
    api=service.spreadsheets()
    title=ANNUAL_TITLE
    meta=api.get(spreadsheetId=sid).execute()
    existing=next((s for s in meta.get("sheets",[]) if s["properties"]["title"]==title),None)
    if existing:
        sheet_id=existing["properties"]["sheetId"]
    else:
        reply=api.batchUpdate(spreadsheetId=sid,body={"requests":[{"addSheet":{"properties":{"title":title}}}]}).execute()
        sheet_id=reply["replies"][0]["addSheet"]["properties"]["sheetId"]
    api.batchUpdate(spreadsheetId=sid,body={"requests":
        d.google_requests(workbook[title],sheet_id,existing)+google_chart_requests(workbook[title],sheet_id)}).execute()


def google_chart_requests(ws, sheet_id):
    """Mirror range-backed FY share charts in native Sheets, preserving missing gaps."""
    from openpyxl.utils.cell import range_boundaries
    requests=[]
    for chart in ws._charts:
        def source_range(formula):
            a,b,c,e=range_boundaries(formula.rsplit("!",1)[1])
            return {"sheetId":sheet_id,"startColumnIndex":a-1,"endColumnIndex":c,
                    "startRowIndex":b-1,"endRowIndex":e}
        first=chart.series[0]
        category=first.cat.numRef.f if first.cat.numRef else first.cat.strRef.f
        title="".join(run.t for paragraph in chart.title.tx.rich.p for run in paragraph.r)
        series=[]
        for i,s in enumerate(chart.series):
            values=source_range(s.val.numRef.f)
            values["startRowIndex"]-=1
            series.append({"series":{"sourceRange":{"sources":[values]}},
                           "targetAxis":"LEFT_AXIS","color":d._rgb("D76A36" if i==0 else "287B59")})
        # Values are fractional shares. Underlying K:L percent formats are kept by the grid serializer.
        domain=source_range(category)
        domain["startRowIndex"]-=1
        requests.append({"addChart":{"chart":{"spec":{"title":title,
            "basicChart":{"chartType":"LINE","legendPosition":"BOTTOM_LEGEND","headerCount":1,
                "interpolateNulls":False,
                "domains":[{"domain":{"sourceRange":{"sources":[domain]}}}],
                "series":series,"axis":[{"position":"BOTTOM_AXIS","title":"Month"},
                    {"position":"LEFT_AXIS","title":"Monthly population share"}]}},
            "position":{"overlayPosition":{"anchorCell":{"sheetId":sheet_id,"rowIndex":chart.anchor._from.row,
                "columnIndex":chart.anchor._from.col},"widthPixels":1000,"heightPixels":260}}}}})
    return requests


def attach_monthly(wb, session, depot, region, selected, folder):
    from src.integrations.google_drive import find_file
    from src.integrations.google_sheets import sheets_service
    existing=find_file(folder,f"{depot}_{selected}")
    cache=read_google_cache(sheets_service(),existing["id"],depot) if existing else new_cache(depot)
    update(cache,[selected],SourceAdapter(session,depot,region),selected,include_upto=True)
    render_tab(wb,cache,selected)
    write_xlsx_cache(wb,cache)
    return cache
