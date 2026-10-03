"""Live source/history and isolated preview checks. Never publish production reports.

APSRTC operations are reads. The optional Sheets check creates two private review
files with only the new tabs, never uses a production file as a write destination,
and never calls Hub, Telegram or Apps Script.
"""
import argparse
import copy
import hashlib
import json
import sys
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

import report_details as d
from src.auth.client import login
from annual_kpi_runner_v3 import tyre_site_info

EXPECTED = {"PRODDUTUR": {"2026-05": (4.94, 4.98), "2026-07": (4.89, 4.94)},
            "RAJAMPET": {"2026-05": (5.20, 5.23), "2026-07": (5.20, 5.22)}}


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()


def totals(cache, period):
    return tuple(cache["months"][period][g]["rows"][-1][-1] for g in ("UD", "UM"))


def existing_annual_model(depot, selected):
    """Read existing KPI data; do not fetch or repair the legacy KPI sources."""
    import annual_history as history
    import annual_kpi_runner_v11 as runner
    folder = "1uQlJJcrv7TbKAXtkFtrrCqZac31aWqL9"
    item = runner.m.find_file(folder, f"{depot}_ANNUAL_KPI_DASHBOARD")
    if not item:
        raise RuntimeError(f"No existing annual KPI file for {depot}; refusing invented core data")
    api = runner.m.sheets_service().spreadsheets()
    meta = api.get(spreadsheetId=item["id"], fields="sheets.properties").execute()
    titles = {s["properties"]["title"] for s in meta["sheets"]}
    if history.CACHE_TITLE in titles:
        cache = history.decode(runner.m.read_values(item["id"], f"'{history.CACHE_TITLE}'!A:D"), depot)
    else:
        title = runner.DETAIL_TITLE if runner.DETAIL_TITLE in titles else runner.m.SHEET_TITLE
        rows = runner.m.read_values(item["id"], f"'{title}'!A:R")
        previous = history.period_from_heading(rows)
        if not previous and runner.m.META_TITLE in titles:
            pairs = runner.m.read_values(item["id"], f"'{runner.m.META_TITLE}'!A:B")
            previous = dict(r[:2] for r in pairs if len(r) > 1).get("LAST_SELECTED_MONTH", "")
        cache = history.new_cache(depot)
        history.migrate(cache, rows, previous)
    year, month = map(int, selected.split("-"))
    fys = runner.m.fy_triplet(year, month)
    matrix = runner.v7.matrix_with_target(history.view_store(cache, fys, selected, runner.m))
    return runner, fys, matrix, item["id"]


def verify_google(api, sid, workbook, cache):
    d.save_google_cache(api, sid, cache)
    d.publish_google_tabs(api, sid, workbook)
    assert d.read_google_cache(api, sid, cache["depot"]) == cache
    meta = api.spreadsheets().get(spreadsheetId=sid, fields="sheets(properties,merges)").execute()
    for title in d.TAB_TITLES:
        ws = workbook[title]
        sheet = next(s for s in meta["sheets"] if s["properties"]["title"] == title)
        grid = sheet["properties"]["gridProperties"]
        assert grid["rowCount"] == ws.max_row + 2 and grid["columnCount"] == ws.max_column
        assert len(sheet.get("merges", [])) == len(ws.merged_cells.ranges)
        rows = api.spreadsheets().values().get(spreadsheetId=sid,
            range=f"'{title}'!A1:{get_column_letter(ws.max_column)}{ws.max_row}",
            valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])
        for r, original in enumerate(ws.values):
            for c, value in enumerate(original):
                got = rows[r][c] if r < len(rows) and c < len(rows[r]) else None
                if r == 2 and c == 0:
                    assert got == "APSRTC"  # native branding; the XLSX has the embedded image
                elif value not in (None, ""):
                    assert got == value, (title, r + 1, c + 1, value, got)
                else:
                    assert got in (None, ""), (title, r + 1, c + 1, got)
    assert next(s for s in meta["sheets"] if s["properties"]["title"] == d.CACHE_TITLE)["properties"]["hidden"]


def full_monthly_preview(depot, cache, session, output):
    """Run the existing generator, replacing only the delivery boundary and new-tab fetch.

    Its prior-history download and source reads remain active. No Drive write or
    delivery function is called; the monthly upload boundary is intercepted.
    """
    import monthly_vehicle_report as monthly
    captured = {}

    def attach(wb, _session, display, region, selected, folder):
        assert display == depot and selected == "2026-05"
        monthly_cache = d.new_cache(depot)
        d.update(monthly_cache, [selected], d.SourceAdapter(session, depot, region, tyre_site_info),
                 refresh_monthly=True)
        assert all(g in monthly_cache["months"][selected] for g in ("B", "F", "UD"))
        assert monthly_cache["months"][selected]["UD"]["rows"][-1][-1] == EXPECTED[depot][selected][0]
        d.render_tabs(wb, monthly_cache, selected)
        d.write_xlsx_cache(wb, monthly_cache)

    def capture(path, **kwargs):
        captured["path"] = Path(path)
        return {"id": "local-review", "webViewLink": "local-review.xlsx"}

    with patch.object(monthly, "login", return_value=session), \
            patch.object(monthly, "attach_monthly_tabs", side_effect=attach), \
            patch.object(monthly, "upload_xlsx_as_google_sheet", side_effect=capture), \
            patch.object(sys, "argv", ["monthly_vehicle_report.py", "--depot", depot, "--month", "2026-05"]):
        assert monthly.main() == 0
    wb = load_workbook(captured["path"])
    assert {"Monthly KMPL", "Vehicle Performance", "Vehicle 360", *d.TAB_TITLES, d.CACHE_TITLE} <= set(wb.sheetnames)
    destination = output / f"{depot}_MONTHLY_2026-05.xlsx"
    wb.save(destination)
    return str(destination)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="reports/detail-review")
    parser.add_argument("--google-previews", action="store_true")
    parser.add_argument("--full-reports", action="store_true")
    args = parser.parse_args()
    out = Path(args.output)
    out.mkdir(parents=True, exist_ok=True)
    session = login()
    result = []
    for depot in EXPECTED:
        cache, wb = d.new_cache(depot), Workbook()
        wb.remove(wb.active)
        adapter = d.SourceAdapter(session, depot, "YSRKADAPA", tyre_site_info)
        api = sid = None
        if args.google_previews:
            from src.integrations.google_sheets import sheets_service
            api = sheets_service()
            created = api.spreadsheets().create(body={"properties": {
                "title": f"REVIEW — {depot} Monthly & Annual Detail Tabs — 2026-10-03"},
                "sheets": [{"properties": {"title": title}} for title in d.TAB_TITLES]}).execute()
            sid = created["spreadsheetId"]  # the only Sheets write destination in this script
        closed = july = None
        steps = []
        for attempt, selected in enumerate(("2026-05", "2026-07", "2026-05"), 1):
            previous = copy.deepcopy(cache["months"])
            calls = d.update(cache, d.annual_periods(selected), adapter, selected_um=selected)
            assert totals(cache, selected) == EXPECTED[depot][selected]
            for period, groups in previous.items():
                for group, snapshot in groups.items():
                    if group != "UM" and not snapshot.get("provisional"):
                        assert cache["months"][period][group] == snapshot
                        assert (group, period) not in calls
            historical = {p: g for p, g in cache["months"].items() if p < "2026-04"}
            if closed is None:
                closed = digest(historical)
            assert digest(historical) == closed
            # Checksum roundtrip is a fresh process's cache input contract.
            cache = d.decode(d.encode(cache), depot)
            d.render_tabs(wb, cache, selected, annual=True)
            d.write_xlsx_cache(wb, cache)
            engine_ws = wb[d.ENGINE_TITLE]
            if attempt == 2:
                july = (engine_ws.max_row, engine_ws.max_column, digest(cache["months"]["2026-07"]))
            if attempt == 3:
                assert engine_ws.max_row < july[0]
                assert digest(cache["months"]["2026-07"]) == july[2]
                visible = d.sections(cache, selected, True)[d.ENGINE_TITLE]
                assert all("July 2026" not in block["title"] for block in visible)
                assert visible[0]["rows"] == cache["months"][selected]["UM"]["rows"]
            if api:
                verify_google(api, sid, wb, cache)
            review_path = out / f"{depot}_ANNUAL_DETAILS_{selected}_step{attempt}.xlsx"
            wb.save(review_path)
            (out / f"{depot}_detail-history_step{attempt}.json").write_text(json.dumps(cache, indent=2))
            record = dict(attempt=attempt, selected=selected, calls=calls,
                          source_totals=dict(zip(("UD", "UM"), totals(cache, selected))),
                          rows=engine_ws.max_row, columns=engine_ws.max_column,
                          fy2526_sha256=closed, errors=cache["errors"])
            steps.append(record)
            print(json.dumps(dict(depot=depot, **record)), flush=True)
        previews = []
        if args.full_reports:
            runner, fys, matrix, original_sid = existing_annual_model(depot, "2026-05")
            runner.REPORT_MONTH = "2026-05"
            path = runner.make_xlsx_v11(depot, matrix, fys, cache)
            target = out / f"{depot}_ANNUAL_KPI_2026-05.xlsx"
            load_workbook(path).save(target)
            previews.append(str(target))
            previews.append(full_monthly_preview(depot, cache, session, out))
        result.append(dict(depot=depot, steps=steps, workbooks=previews,
                           google_preview=f"https://docs.google.com/spreadsheets/d/{sid}/edit" if sid else None,
                           production_writes=0, deliveries=0))
        (out / "validation-summary.json").write_text(json.dumps(result, indent=2))
    print("DETAIL_REVIEW_SUCCESS", flush=True)


if __name__ == "__main__":
    main()
