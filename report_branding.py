"""Presentation-only helpers. Never change APSRTC values or history."""
import math
import re
from pathlib import Path

LOGO_URL = 'https://www.apsrtc.ap.gov.in/images/apsrtc_logo1.gif'
LOGO_PATH = Path(__file__).resolve().parent / 'assets' / 'apsrtc_logo.png'


def kmpl_number(value):
    if value is None or isinstance(value, bool):
        return None
    match = re.fullmatch(r'\s*([+-]?\d+(?:\.\d+)?)\s*(?:(?:🔧|⚙)[\s\d]*)?\s*', str(value))
    if not match:
        return None
    result = float(match.group(1))
    return result if math.isfinite(result) else None


def style_kmpl(cell):
    """Keep annotations and zeros intact; grade using existing KMPL bands."""
    from openpyxl.styles import PatternFill, Font
    value = kmpl_number(cell.value)
    if value is None:
        return
    bands = [(5, 'F4CCCC', '9C0006'), (5.1, 'FCE5CD', '7F6000'),
             (5.2, 'FFF2CC', '7F6000'), (5.3, 'D9EAD3', '274E13')]
    fill, text = next(((f, t) for limit, f, t in bands if value <= limit), ('B6D7A8', '274E13'))
    cell.fill = PatternFill('solid', fgColor=fill)
    cell.font = Font(name='Arial', size=10, color=text, bold=value <= 5 or value > 5.3)


def add_logo(ws, anchor='A1', width=180):
    from openpyxl.drawing.image import Image
    image = Image(str(LOGO_PATH))
    image.width = width
    image.height = width * 130 / 640
    ws.add_image(image, anchor)


def style_monthly_workbook(wb):
    for name, first_row, first_col, last_col in [('Monthly KMPL', 6, 5, None),
            ('Vehicle Performance', 6, 6, 17), ('Vehicle360', 5, 5, 5),
            ('Vehicle 360', 5, 5, 5)]:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if not ws._images:
            if name == 'Monthly KMPL':
                from openpyxl.utils import get_column_letter
                width = sum(ws.column_dimensions[get_column_letter(c)].width * 7 + 5 for c in (1, 2, 3)) - 10
                add_logo(ws, anchor='A3', width=width)
                ws.row_dimensions[3].height = math.ceil(width * 130 / 640 * .75) + 3
            else:
                add_logo(ws)
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows(min_row=first_row, min_col=first_col, max_col=last_col or ws.max_column):
            # History footer / exceptions are not vehicle performance records.
            if name == 'Vehicle Performance' and str(ws.cell(row[0].row, 5).value or '') not in {'2024-25', '2025-26', '2026-27'}:
                continue
            for cell in row:
                style_kmpl(cell)



def add_monthly_management_dashboard(wb, display, report_period, vehicle_data):
    """Create a decision-first front page from values already fetched for the Monthly report."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter

    title = "1. Management Dashboard"
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, 0)
    ws.sheet_view.showGridLines = False
    navy, blue, green, red, amber, pale, ink = "12345B","1976B9","287B59","A33232","B66521","F3F6FA","23364D"
    line = Side(style="thin", color="DCE5EF")

    vehicles = []
    total_km = total_hsd = 0.0
    for vehicle, data in vehicle_data.items():
        km = float(data.get("total_kms") or 0)
        hsd = float(data.get("total_hsd") or 0)
        kmpl = data.get("up_to_day_latest")
        try: kmpl = float(kmpl) if kmpl not in (None, "") else None
        except (TypeError, ValueError): kmpl = None
        total_km += km
        total_hsd += hsd
        vehicles.append((vehicle, kmpl, data.get("op_type",""), data.get("engine","")))
    depot_kmpl = total_km / total_hsd if total_hsd > 0 else None
    measured = [v for v in vehicles if v[1] is not None]
    poor = sorted([v for v in measured if v[1] <= 5.00], key=lambda x:(x[1],x[0]))
    strong = [v for v in measured if v[1] > 5.30]

    # Driver/vehicle range counts are read only from the already-rendered source-backed table.
    range_counts = {}
    if "KMPL Range Distribution" in wb.sheetnames:
        rs = wb["KMPL Range Distribution"]
        header_row = next((r for r in range(1, min(rs.max_row,20)+1)
                           if str(rs.cell(r,1).value or "").strip() == "KMPL Range"), None)
        if header_row:
            headers = {str(rs.cell(header_row,c).value or "").strip():c for c in range(1,rs.max_column+1)}
            for r in range(header_row+1, rs.max_row+1):
                label = str(rs.cell(r,1).value or "").strip()
                if not label: continue
                for key in ("Vehicle Month","Vehicle Upto","Driver Month","Driver Upto"):
                    c=headers.get(key)
                    if c:
                        range_counts[(label,key)] = rs.cell(r,c).value

    for c,w in enumerate([16,18,18,18,18,18,18,18,18,18,18,18],1):
        ws.column_dimensions[get_column_letter(c)].width=w
    ws.merge_cells("A1:L1"); ws["A1"]=f"APSRTC — {display} DEPOT — MONTHLY MANAGEMENT DASHBOARD"
    ws["A1"].fill=PatternFill("solid",fgColor=navy); ws["A1"].font=Font(name="Arial",size=18,bold=True,color="FFFFFF"); ws["A1"].alignment=Alignment(horizontal="center")
    ws.merge_cells("A2:L2"); ws["A2"]=f"{report_period} | Decision view — detailed source sheets remain unchanged"
    ws["A2"].font=Font(name="Arial",size=11,bold=True,color=ink); ws["A2"].alignment=Alignment(horizontal="center")
    cards=[("DEPOT KMPL",depot_kmpl,"0.00",blue),("TOTAL KM",total_km,"#,##0",navy),
           ("HSD USED (L)",total_hsd,"#,##0.0",navy),("POOR VEHICLES ≤5.00",len(poor),"0",red),
           ("STRONG VEHICLES >5.30",len(strong),"0",green)]
    for i,(label,value,fmt,color) in enumerate(cards):
        col=1+i*2
        ws.merge_cells(start_row=4,start_column=col,end_row=4,end_column=col+1)
        ws.merge_cells(start_row=5,start_column=col,end_row=6,end_column=col+1)
        a=ws.cell(4,col,label); a.fill=PatternFill("solid",fgColor=color); a.font=Font(name="Arial",size=9,bold=True,color="FFFFFF"); a.alignment=Alignment(horizontal="center")
        v=ws.cell(5,col,"—" if value is None else value); v.fill=PatternFill("solid",fgColor=pale); v.font=Font(name="Arial",size=18,bold=True,color=ink); v.alignment=Alignment(horizontal="center",vertical="center"); v.number_format=fmt

    ws.merge_cells("A8:F8"); ws["A8"]="Management Attention"; ws["A8"].fill=PatternFill("solid",fgColor=red); ws["A8"].font=Font(name="Arial",size=12,bold=True,color="FFFFFF")
    ws.merge_cells("G8:L8"); ws["G8"]="Positive Signals"; ws["G8"].fill=PatternFill("solid",fgColor=green); ws["G8"].font=Font(name="Arial",size=12,bold=True,color="FFFFFF")
    attention = []
    if poor: attention.append(f"{len(poor)} vehicles are at or below 5.00 KMPL; review the lowest performers first.")
    if depot_kmpl is not None and depot_kmpl <= 5.10: attention.append(f"Depot weighted KMPL is {depot_kmpl:.2f}; fuel-efficiency attention is required.")
    if not attention: attention.append("No current-month exception crossed the dashboard attention rules.")
    positive=[]
    if strong: positive.append(f"{len(strong)} vehicles are above 5.30 KMPL; preserve and study the operating pattern.")
    if depot_kmpl is not None and depot_kmpl > 5.20: positive.append(f"Depot weighted KMPL is {depot_kmpl:.2f}, above the 5.20 dashboard watch level.")
    if not positive: positive.append("No positive rule triggered; use detailed tabs for neutral/source context.")
    for i,textv in enumerate(attention[:3],9):
        ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=6); ws.cell(i,1,textv).alignment=Alignment(wrap_text=True,vertical="top"); ws.cell(i,1).font=Font(name="Arial",size=10,color=ink)
    for i,textv in enumerate(positive[:3],9):
        ws.merge_cells(start_row=i,start_column=7,end_row=i,end_column=12); ws.cell(i,7,textv).alignment=Alignment(wrap_text=True,vertical="top"); ws.cell(i,7).font=Font(name="Arial",size=10,color=ink)

    ws.merge_cells("A13:F13"); ws["A13"]="Top 5 Problem Vehicles"; ws["A13"].fill=PatternFill("solid",fgColor=navy); ws["A13"].font=Font(name="Arial",bold=True,color="FFFFFF")
    for c,h in enumerate(("Vehicle","Month-end KMPL","Operation","Engine"),1): ws.cell(14,c,h)
    for r,v in enumerate(poor[:5],15):
        for c,val in enumerate((v[0],v[1],v[2],v[3]),1): ws.cell(r,c,val)
        ws.cell(r,2).number_format="0.00"
    ws.merge_cells("G13:L13"); ws["G13"]="Driver / Vehicle Range Focus"; ws["G13"].fill=PatternFill("solid",fgColor=navy); ws["G13"].font=Font(name="Arial",bold=True,color="FFFFFF")
    ws["G14"]="Source range"; ws["H14"]="Vehicle Month"; ws["I14"]="Vehicle Upto"; ws["J14"]="Driver Month"; ws["K14"]="Driver Upto"
    focus_labels=["Below -3.00","3.00 - 4.00","4.01 - 5.00"]
    for r,label in enumerate(focus_labels,15):
        ws.cell(r,7,label)
        for c,key in enumerate(("Vehicle Month","Vehicle Upto","Driver Month","Driver Upto"),8):
            ws.cell(r,c,range_counts.get((label,key),"—"))

    ws.merge_cells("A22:L22"); ws["A22"]="Key Findings & Recommended Actions"; ws["A22"].fill=PatternFill("solid",fgColor=navy); ws["A22"].font=Font(name="Arial",size=12,bold=True,color="FFFFFF")
    findings=[]
    if poor:
        names=", ".join(v[0] for v in poor[:5])
        findings.append(f"Priority engineering/driver review: {names}. These are the lowest current month-end KMPL vehicles at or below 5.00.")
    if total_hsd>0:
        findings.append(f"Fuel context: {total_km:,.0f} km used {total_hsd:,.1f} litres, giving weighted depot KMPL {depot_kmpl:.2f}. Use vehicle detail to locate avoidable consumption.")
    driver_poor=sum((range_counts.get((label,"Driver Month")) or 0) for label in focus_labels if isinstance(range_counts.get((label,"Driver Month")), (int,float)))
    if driver_poor:
        findings.append(f"Driver focus: {int(driver_poor)} Month-scope driver observations are in the three lowest source ranges shown above; prioritise coaching from the detailed range sheet.")
    findings.append("Tyre/maintenance decisions should use the existing B/C/F and Vehicle 360 detail tabs; this dashboard intentionally surfaces only exceptions supported by current source data.")
    for i,textv in enumerate(findings[:4],23):
        ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=12); ws.cell(i,1,textv); ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top"); ws.cell(i,1).font=Font(name="Arial",size=10,color=ink); ws.row_dimensions[i].height=30

    for row in ws.iter_rows(min_row=14,max_row=19,min_col=1,max_col=11):
        for cell in row:
            cell.border=Border(bottom=line); cell.font=Font(name="Arial",size=9,bold=cell.row==14); cell.alignment=Alignment(vertical="center",wrap_text=True)
    add_logo(ws,"J1",width=190)
    ws.freeze_panes="A3"
    ws.print_area="A1:L27"; ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws
