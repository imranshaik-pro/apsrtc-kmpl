"""Owner template, source C hierarchy and real calculation dependency regressions."""
import copy
import unittest
from pathlib import Path
from unittest.mock import Mock
from openpyxl import Workbook
import report_details as d

FIXTURES = Path(__file__).parent / "fixtures/report-details"


def snapshot(group, period="2026-05", depot="PRODDUTUR"):
    code = "PDT" if depot == "PRODDUTUR" else "RJP"
    return d.parse_tyre((FIXTURES/depot/f"{period}-{group}.html").read_text(), group, code, period, depot)


def cache():
    result = d.new_cache("PRODDUTUR")
    result["months"]["2026-05"] = {g:snapshot(g) for g in ("B", "C", "F")}
    return result


class TyreTemplateTests(unittest.TestCase):
    def test_c_groups_have_unique_keys_exact_totals_and_distinct_stage_codes(self):
        for depot, expected in (("PRODDUTUR", [3,13,12,2]), ("RAJAMPET", [0,2,5,4])):
            snap = snapshot("C", depot=depot)
            self.assertEqual(len(snap["headers"]), 47)
            self.assertEqual(len(set(map(d.key,snap["headers"]))),47)
            self.assertEqual([snap["rows"][0][i] for i in (16,26,36,46)], expected)
            self.assertEqual(snap["rows"][0][5],sum(expected))
        snap = snapshot("C")
        self.assertEqual(snap["rows"][0][7],1)
        self.assertEqual(snap["rows"][0][17],4)

    def test_c_wrong_month_depot_header_order_or_duplicate_conflict_is_rejected(self):
        html=(FIXTURES/"PRODDUTUR/2026-05-C.html").read_text()
        for text, depot, period in ((html,"RJP","2026-05"),(html,"PDT","2026-06"),
                (html.replace('<th>S1</th>','<th>S9</th>',1),"PDT","2026-05")):
            with self.assertRaises(ValueError):
                d.parse_tyre(text,"C",depot,period,"PRODDUTUR")

    def test_c_backfill_does_not_refetch_existing_closed_b_f_or_ud(self):
        saved=cache(); old=copy.deepcopy(saved["months"]["2026-05"])
        del saved["months"]["2026-05"]["C"]
        saved["months"]["2026-05"]["UD"] = dict(depot="PRODDUTUR",period="2026-05",group="UD",headers=["x"],rows=[[0]])
        fetch=Mock(return_value=old["C"])
        self.assertEqual(d.update(saved,["2026-05"],fetch),[("C","2026-05")])
        self.assertEqual(saved["months"]["2026-05"]["B"],old["B"])
        self.assertEqual(saved["months"]["2026-05"]["F"],old["F"])
        self.assertEqual(d.decode(d.encode(saved),"PRODDUTUR"),saved)

    def test_monthly_template_has_only_requested_columns_and_hides_denominators(self):
        wb=Workbook();d.render_tabs(wb,cache(),"2026-05")
        blocks=d.sections(cache(),"2026-05",False)[d.TYRE_TITLE]
        self.assertEqual(len(blocks),3)
        self.assertEqual(blocks[0]["headers"][-3:],["Mech. Defects %","Stone %","Worn Smooth %"])
        self.assertEqual(blocks[0]["rows"][0][-3:],[.375,0,.125])
        self.assertEqual(blocks[1]["rows"][0][-1],7)
        self.assertEqual(blocks[2]["rows"][0][21],13)
        all_values=[cell.value for row in wb[d.TYRE_TITLE] for cell in row]
        for label in ("Zone Name","Region","Depot","Number of Tyres Received","Repair","Dcp_Code"):
            self.assertNotIn(label,all_values)
        self.assertEqual(wb[d.TYRE_INPUT_TITLE].sheet_state,"hidden")
        self.assertEqual(list(wb[d.TYRE_INPUT_TITLE].values)[1][1:],(32,73,0,4,0,32,0))
        self.assertTrue(any("source total 13; S1–S9 sum 10" in str(v) for v in all_values))

    def test_annual_formulas_sum_counts_and_use_distinct_matching_denominators(self):
        wb=Workbook();d.render_tabs(wb,cache(),"2026-05",annual=True)
        ws=wb[d.TYRE_TITLE]
        totals=[r for r in range(1,ws.max_row+1) if ws.cell(r,2).value=="FY Total"]
        self.assertEqual(len(totals),6)
        for r in totals[:2]:
            self.assertEqual(ws.cell(r,11).data_type,"f")
            self.assertIn("SUM(K",ws.cell(r,12).value)
            self.assertIn("'_TYRE_INPUTS'!B",ws.cell(r,12).value)
            self.assertIn("'_TYRE_INPUTS'!C",ws.cell(r,13).value)
            self.assertIn("'_TYRE_INPUTS'!D",ws.cell(r,13).value)
            self.assertIn("'_TYRE_INPUTS'!G",ws.cell(r,14).value)
            self.assertIn("COUNT(",ws.cell(r,12).value)
            self.assertEqual(ws.cell(r,12).number_format,"0.00%")
        requests=d.google_requests(ws,123)
        cells=next(x["updateCells"]["rows"] for x in requests if "updateCells" in x)
        r=totals[0]
        self.assertEqual(cells[r-1]["values"][11]["userEnteredValue"],{"formulaValue":ws.cell(r,12).value})
        self.assertEqual(cells[r-1]["values"][11]["userEnteredFormat"]["numberFormat"]["type"],"PERCENT")

    def test_july_to_may_removes_future_tyre_rows_and_formula_input_ranges(self):
        saved=cache();saved["months"]["2026-07"]={"C":snapshot("C","2026-07")}
        wb=Workbook();d.render_tabs(wb,saved,"2026-07",annual=True)
        before=wb[d.TYRE_TITLE].max_row
        d.render_tabs(wb,saved,"2026-05",annual=True)
        self.assertLess(wb[d.TYRE_TITLE].max_row,before)
        self.assertFalse(any(c.value=="Jul-26" for row in wb[d.TYRE_TITLE] for c in row))
        self.assertEqual(wb[d.TYRE_INPUT_TITLE].max_row,15)
        self.assertEqual(saved["months"]["2026-07"]["C"]["rows"][0][5],27)


if __name__ == "__main__":
    unittest.main()
