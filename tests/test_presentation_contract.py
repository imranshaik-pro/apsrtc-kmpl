"""Offline tests: no source login, Drive writes, Telegram messages or dispatches."""
import ast
import importlib.util
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import annual_visuals as visual
import telegram_command_bot as bot


class PresentationTests(unittest.TestCase):
    def test_print_identity_and_monthly_heading(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        tree=ast.parse((ROOT/'monthly_vehicle_report.py').read_text())
        functions=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ('get_style','apply_formatting')]
        class PD:
            @staticmethod
            def isna(v): return False
        thin=Side(style='thin',color='B7C9E2')
        env=dict(pd=PD,Alignment=Alignment,Border=Border,Font=Font,PatternFill=PatternFill,Side=Side,
                 get_column_letter=get_column_letter,THIN=thin,THIN_BORDER=Border(left=thin,right=thin,top=thin,bottom=thin))
        exec(compile(ast.Module(body=functions,type_ignores=[]),'<monthly>','exec'),env)
        w=Workbook(); s=w.active; s.title='Monthly KMPL'
        s.append(['SL No','Vehicle No','Op Type','Engine Type','1','Up-To-Day'])
        s.append([1,'04Z1234','EXPRESS','TATA',5.2,5.3])
        env['apply_formatting'](w,'RAJAMPET','June 2026')
        self.assertIn('RAJAMPET',s['A1'].value)
        self.assertIn('June 2026',s['A2'].value)
        self.assertEqual(s.print_title_rows,'$1:$5')
        self.assertIn('June 2026',s.oddHeader.center.text)

    def test_history_reader_accepts_identity_rows(self):
        from openpyxl import Workbook
        import re
        tree=ast.parse((ROOT/'src/reporting/vehicle_history.py').read_text())
        functions=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ('norm_vehicle','read_existing_history')]
        env=dict(re=re,FYS=('2024-25','2025-26','2026-27'),MONTHS=('Apr',))
        exec(compile(ast.Module(body=functions,type_ignores=[]),'<history>','exec'),env)
        w=Workbook(); s=w.active
        s.append(['APSRTC RAJAMPET']); s.append(['June 2026']); s.append(['Veh No','FY Year','Apr'])
        s.append(['AP04Z1234','2026-27',5.2])
        self.assertEqual(env['read_existing_history'](s)['04Z1234']['2026-27']['Apr'],5.2)

    def test_ragged_google_rows_retain_hsd_label_and_zero(self):
        rows=[["SL.No","KPI","Year"],[1,"HSD KMPL EXCL AC","2024-25",5.13],
              ["","","2025-26"]+[""]*14+[5.24],
              ["","","2026-27"]+[""]*14+[0]]
        years=["2024-25","2025-26","2026-27"]
        data=visual.records(rows,years)
        self.assertEqual(data['HSD KMPL EXCL AC']['2026-27'][17],0)
        model=visual.dashboard_model('RAJAMPET',years,rows,'2026-06')
        self.assertTrue(any(c['text']==0 and c['size']==24 for c in model['cells']))
        self.assertEqual([c['points'] for c in model['charts']],[3,3])
        self.assertTrue(any('June 2026' in str(c['text']) for c in model['cells']))
        self.assertEqual(len([r for r in visual.google_dashboard_requests(12,model) if 'addChart' in r]),2)

    def test_month_boundaries_and_unavailable_values(self):
        for month,expected in [('2026-04',1),('2026-08',5),('2027-03',12)]:
            model=visual.dashboard_model('RAJAMPET',['2024-25','2025-26','2026-27'],[],month)
            self.assertEqual(model['charts'][0]['points'],expected)
        self.assertIsNone(visual.number('MANUAL INPUT REQUIRED'))
        self.assertEqual(visual.number('0.00'),0)

    def test_bot_annual_dispatch_and_future_rejection(self):
        with patch.object(bot,'DEPOTS',('RAJAMPET',)), patch.object(bot,'send'), patch.object(bot,'answer_callback'), patch.object(bot,'dispatch') as dispatch:
            bot.handle_callback('token','chat',{'id':'cb','data':'annualrun|RAJAMPET|2026-06'})
            dispatch.assert_called_once_with('annual-kpi.yml',{'depot':'RAJAMPET','selected_month':'2026-06','financial_years':'2026-27'})
            dispatch.reset_mock()
            bot.handle_callback('token','chat',{'id':'cb','data':'annualrun|RAJAMPET|2099-06'})
            dispatch.assert_not_called()

    def test_vehicle_event_input_and_confirmation(self):
        self.assertEqual(bot.normalize_vehicle('ap 04 z 1234'),'04Z1234')
        self.assertEqual(bot.normalize_vehicle('04Z1234'),'04Z1234')
        self.assertEqual(bot.normalize_vehicle('invalid'),'')
        prompt=bot.event_detail_prompt('RAJAMPET','AP04Z1234','BREAKDOWN','2026-06-22')
        payload=bot.parse_event_details(prompt,'Bus stand | 25 | Fuel pump | Repaired')
        self.assertEqual(payload['kms_cancelled'],'25')
        self.assertIn('Reply CONFIRM',bot.confirmation_text(payload))

    def test_daily_monthly_dispatch(self):
        with patch.object(bot,'DEPOTS',('RAJAMPET',)), patch.object(bot,'send'), patch.object(bot,'answer_callback'), patch.object(bot,'dispatch') as dispatch:
            bot.handle_callback('token','chat',{'data':'dailyrun|RAJAMPET|latest'})
            dispatch.assert_called_with('daily-report.yml',{'depot':'RAJAMPET','report_date':''})
            bot.handle_callback('token','chat',{'data':'monthlyrun|RAJAMPET|2026-06'})
            dispatch.assert_called_with('monthly-report.yml',{'depot':'RAJAMPET','month':'2026-06'})

if __name__=='__main__': unittest.main()
