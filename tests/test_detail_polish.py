"""Source-preserving layout and both-depot selection regressions."""
import ast
import copy
import unittest
from pathlib import Path
from openpyxl import Workbook
import report_details as d
from tests.test_report_details import fixture, engine, FIXTURES
from tests.test_tyre_template import cache as tyre_cache

class DetailPolishTests(unittest.TestCase):
    def test_every_statement_c_field_survives_stage_panels(self):
        cache=tyre_cache(); before=copy.deepcopy(cache)
        original=d.sections(cache,'2026-05',False)[d.TYRE_TITLE][-1]
        panels=[b for b in d.presentation_sections(cache,'2026-05',False)[d.TYRE_TITLE] if b.get('template_kind')=='scrap']
        self.assertEqual(len(panels),4)
        self.assertEqual([v for p in panels for v in p['rows'][0][2:]], original['rows'][0][2:])
        self.assertEqual([p['title'].split(' · ')[-1] for p in panels],list(d.SCRAP_GROUPS))
        self.assertEqual(cache,before)
        wb=Workbook();d.render_tabs(wb,cache,'2026-05')
        self.assertEqual(wb[d.TYRE_TITLE].max_column,14)
        self.assertEqual(wb[d.TYRE_TITLE].freeze_panes,'A5')
        self.assertTrue(any('source total 13; S1–S9 sum 10' in str(c.value) for row in wb[d.TYRE_TITLE] for c in row))

    def test_long_dynamic_matrix_has_each_row_once_and_only_one_total(self):
        cache=d.new_cache('RAJAMPET');snap=engine('RAJAMPET','2026-05','UD')
        rows=[[i,'Unique engine '+str(i)]+[5.123456]*(len(snap['headers'])-2) for i in range(1,46)]
        rows.append(['TOTAL','']+[5.2]*(len(snap['headers'])-2));snap['rows']=rows
        cache['months']['2026-05']={'UD':snap};before=copy.deepcopy(cache)
        panels=d.presentation_sections(cache,'2026-05',False)[d.ENGINE_TITLE]
        self.assertEqual([r for b in panels for r in b['rows']],rows)
        self.assertTrue(all(len(b['rows'])<=20 for b in panels))
        wb=Workbook();d.render_tabs(wb,cache,'2026-05')
        values=[c.value for row in wb[d.ENGINE_TITLE] for c in row]
        for i in range(1,46):self.assertEqual(values.count('Unique engine '+str(i)),1)
        self.assertEqual(values.count('TOTAL'),1)
        self.assertEqual(cache,before)

    def test_both_depots_may_july_may_retain_history_and_remove_future_categories(self):
        for depot in ('PRODDUTUR','RAJAMPET'):
            with self.subTest(depot=depot):
                cache=d.new_cache(depot);fetch=lambda g,p: fixture(g,p,depot)
                d.update(cache,['2026-04','2026-05'],fetch,selected_um='2026-05')
                closed=copy.deepcopy(cache['months']);wb=Workbook()
                d.render_tabs(wb,cache,'2026-05',annual=True)
                initial=[[c.value for c in row] for row in wb[d.ENGINE_TITLE]]
                d.update(cache,['2026-04','2026-05','2026-06','2026-07'],fetch,selected_um='2026-07')
                d.render_tabs(wb,cache,'2026-07',annual=True)
                july_rows=wb[d.ENGINE_TITLE].max_row
                saved=copy.deepcopy(cache)
                d.update(cache,['2026-04','2026-05'],fetch,selected_um='2026-05')
                d.render_tabs(wb,cache,'2026-05',annual=True)
                self.assertEqual(initial,[[c.value for c in row] for row in wb[d.ENGINE_TITLE]])
                self.assertLess(wb[d.ENGINE_TITLE].max_row,july_rows)
                for p in closed:
                    for g in ('B','C','F','UD'):self.assertEqual(cache['months'][p][g],closed[p][g])
                self.assertEqual(cache['months']['2026-07'],saved['months']['2026-07'])
                self.assertFalse(any('July 2026' in str(c.value) for row in wb[d.ENGINE_TITLE] for c in row))
                req=d.google_requests(wb[d.ENGINE_TITLE],17)
                grid=next(r['updateCells'] for r in req if 'updateCells' in r)
                self.assertEqual(grid['range']['endRowIndex'],wb[d.ENGINE_TITLE].max_row+2)
                self.assertTrue(any(v.get('note') and 'separate endpoints' in v['note'] for r in grid['rows'] for v in r['values']))

    def test_monthly_formatting_preserves_precision_and_annotations(self):
        # Load only presentation functions; no source login or integration imports.
        from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
        from openpyxl.utils import get_column_letter
        import pandas as pd
        tree=ast.parse(Path('monthly_vehicle_report.py').read_text())
        fs=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ('get_style','apply_formatting')]
        thin=Side(style='thin',color='B7C9E2')
        env=dict(pd=pd,Alignment=Alignment,Border=Border,Font=Font,PatternFill=PatternFill,Side=Side,get_column_letter=get_column_letter,THIN=thin,THIN_BORDER=Border(left=thin,right=thin,top=thin,bottom=thin))
        exec(compile(ast.Module(body=fs,type_ignores=[]),'<monthly presentation>','exec'),env)
        wb=Workbook();ws=wb.active;ws.title='Monthly KMPL'
        ws.append(['SNo','Vehicle','Engine','Product','1','2','UD']);ws.append([1,'04Z1234','TATA','EX',5.123456,'4.52 🔧',0])
        env['apply_formatting'](wb,'PRODDUTUR','May 2026')
        from report_branding import style_monthly_workbook
        style_monthly_workbook(wb);style_monthly_workbook(wb)
        self.assertEqual([ws.cell(6,c).value for c in (5,6,7)],[5.123456,'4.52 🔧',0])
        self.assertEqual(ws['E6'].number_format,'0.00')
        self.assertEqual(len(ws._images),1)
        self.assertEqual(ws._images[0].anchor,'A3')
        self.assertIn('A3:C3',[str(r) for r in ws.merged_cells.ranges])

if __name__=='__main__':unittest.main()
