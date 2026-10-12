"""Offline presentation regressions; no report publication or source login."""
import unittest
from io import BytesIO
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from report_branding import kmpl_number, style_monthly_workbook, add_monthly_management_dashboard, LOGO_URL
from annual_visuals import dashboard_model, google_dashboard_requests, render_xlsx_dashboard


class BrandingTests(unittest.TestCase):
    def test_annotations_missing_and_zero(self):
        for value, expected in [('4.52 🔧 18', 4.52), ('5.31 ⚙', 5.31), (0, 0), ('—', None), ('MANUAL INPUT REQUIRED', None), ('', None), (True, None)]:
            self.assertEqual(kmpl_number(value), expected)

    def test_all_monthly_views_preserve_values_and_colour_annotations(self):
        w=Workbook(); w.remove(w.active)
        for name, row, col in [('Monthly KMPL',6,5),('Vehicle Performance',6,6),('Vehicle 360',5,5)]:
            s=w.create_sheet(name); s.cell(row,col,'4.52 🔧 18')
            s.cell(row+1,col,5.4); s.cell(row+2,col,0)
            if name=='Vehicle Performance':
                for r in range(row,row+3): s.cell(r,5,'2024-25')
        original={s.title:{c.coordinate:c.value for row in s for c in row if c.value is not None} for s in w}
        style_monthly_workbook(w)
        for s in w:
            self.assertEqual(original[s.title],{c.coordinate:c.value for row in s for c in row if c.value is not None})
            row,col=(5,5) if s.title=='Vehicle 360' else (6,6) if s.title=='Vehicle Performance' else (6,5)
            self.assertEqual(s.cell(row,col).fill.fgColor.rgb[-6:],'F4CCCC')
            self.assertEqual(s.cell(row+1,col).fill.fgColor.rgb[-6:],'B6D7A8')
            self.assertEqual(s.cell(row+2,col).fill.fgColor.rgb[-6:],'F4CCCC')
        b=BytesIO();w.save(b)
        self.assertEqual(len([p for p in ZipFile(b).namelist() if p.startswith('xl/media/')]),3)
        b.seek(0); saved=load_workbook(b)
        self.assertEqual(saved['Vehicle Performance']['F6'].value,'4.52 🔧 18')

    def test_annual_directions_missing_target_and_google_logo(self):
        fys=['2024-25','2025-26','2026-27']; rows=[]
        for name,target,value in [('TOTAL LUB KMPL',1600,1700),('B.D RATE',.03,.04),('MED CANCL.',.06,.06),('HSD KMPL INCL AC','',5.2)]:
            for fy in fys: rows.append([1,name,fy,target]+['']*13+[value])
        model=dashboard_model('PRODDUTUR',fys,rows,'2026-06')
        for name,fill in [('TOTAL LUB KMPL','E2F0D9'),('B.D RATE','F4CCCC'),('MED CANCL.','E2F0D9'),('HSD KMPL INCL AC','FFF2D9')]:
            r=next(c['r'] for c in model['cells'] if c['text']==name)
            self.assertEqual(next(c['fill'] for c in model['cells'] if c['r']==r and c['c']==7),fill)
        requests=google_dashboard_requests(12,model)
        self.assertIn(LOGO_URL,str(requests));self.assertIn('formulaValue',str(requests))
        w=Workbook();render_xlsx_dashboard(w,'Dashboard',model,'PRODDUTUR','2026-06')
        b=BytesIO();w.save(b)
        self.assertTrue(any(p.startswith('xl/media/') for p in ZipFile(b).namelist()))


    def test_monthly_management_dashboard_is_first_and_uses_source_values(self):
        w=Workbook(); w.active.title="Monthly KMPL"
        r=w.create_sheet("KMPL Range Distribution")
        r.append(["KMPL Range","Vehicle Month","Vehicle Upto","Driver Month","Driver Upto"])
        r.append(["Below -3.00",1,2,3,4]); r.append(["3.00 - 4.00",2,3,4,5]); r.append(["4.01 - 5.00",3,4,5,6])
        vehicles={
            "AP01A1":{"total_kms":500,"total_hsd":100,"up_to_day_latest":5.0,"op_type":"EXP","engine":"E1"},
            "AP01A2":{"total_kms":560,"total_hsd":100,"up_to_day_latest":5.6,"op_type":"ORD","engine":"E2"},
        }
        add_monthly_management_dashboard(w,"PRODDUTUR","September 2026",vehicles)
        self.assertEqual(w.sheetnames[0],"1. Management Dashboard")
        d=w["1. Management Dashboard"]
        self.assertAlmostEqual(d["A5"].value,5.3)
        self.assertEqual(d["G15"].value,"Below -3.00")
        self.assertEqual(d["J15"].value,3)
        self.assertIn("AP01A1",d["A23"].value)

    def test_annual_scorecard_excludes_detail_rows_and_adds_actions(self):
        fys=["2024-25","2025-26","2026-27"]; rows=[]
        for name,target,value in [("HSD KMPL INCL AC",5.3,5.1),("B.D RATE",.03,.04),("PRODUCT: EXPRESS",5.4,5.2),("ENGINE: TEST",5.4,5.1)]:
            for fy in fys: rows.append([1,name,fy,target]+[""]*13+[value])
        model=dashboard_model("PRODDUTUR",fys,rows,"2026-06")
        texts=[c["text"] for c in model["cells"]]
        self.assertIn("Key Findings & Recommended Actions",texts)
        self.assertFalse(any(t=="EXPRESS" for t in texts))
        self.assertFalse(any(t=="TEST" for t in texts))
        self.assertTrue(any(isinstance(t,str) and t.startswith("ATTENTION") for t in texts))

if __name__=='__main__': unittest.main()

