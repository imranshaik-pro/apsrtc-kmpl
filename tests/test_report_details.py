"""Real source contracts, immutable history and generated-view lifecycle regressions."""
import copy
import json
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
import report_details as d

FIXTURES = Path(__file__).parent / "fixtures" / "report-details"


def engine(depot, period, mode):
    return d.parse_engine((FIXTURES / depot / f"{period}-{mode}.html").read_text(),
                          depot, period, mode)


def fixture(group, period, depot="PRODDUTUR"):
    if group in ("UD", "UM") and (FIXTURES / depot / f"{period}-{group}.html").exists():
        return engine(depot, period, group)
    # Synthetic periods exercise persistence, not live-source correctness.
    if group in ("B", "C", "F"):
        return dict(depot=depot, period=period, group=group,
                    headers=["S.N", "Month-Year", "Depot", "Tyre size", "Count"],
                    rows=[[1, d.month_end(period).strftime("%b-%Y").upper(),
                           "PDT", "All Tyre Sizes Total", 0]],
                    through=d.month_end(period).isoformat())
    return dict(depot=depot, period=period, group=group,
                headers=["SNO", "ENGINE MODEL", "EXPRESS", "Grand Total"],
                rows=[[1, "Exact source engine", 0, 0], ["TOTAL", "", 0, 0]],
                start=(d.month_end(period).replace(day=1) if group == "UD" else d.fy_start(period)).isoformat(),
                through=d.month_end(period).isoformat())


class SourceContractTests(unittest.TestCase):
    def test_verified_ud_month_and_um_financial_year_periods(self):
        for depot, expected in (("PRODDUTUR", (4.94, 4.98, 4.89, 4.94)),
                                ("RAJAMPET", (5.20, 5.23, 5.20, 5.22))):
            for (period, mode), value in zip(
                    (("2026-05", "UD"), ("2026-05", "UM"), ("2026-07", "UD"), ("2026-07", "UM")), expected):
                with self.subTest(depot=depot, period=period, mode=mode):
                    snap = engine(depot, period, mode)
                    self.assertEqual(snap["rows"][-1][-1], value)
                    self.assertEqual(snap["start"], f"{period}-01" if mode == "UD" else "2026-04-01")

    def test_total_colspan_keeps_every_product_total_in_its_column(self):
        snap = engine("PRODDUTUR", "2026-05", "UD")
        totals = dict(zip(snap["headers"], snap["rows"][-1]))
        self.assertEqual(totals["SNO"], "TOTAL")
        self.assertEqual(totals["ENGINE MODEL"], "")
        self.assertEqual(totals["CONTAINER"], 4.56)
        self.assertEqual(totals["DGT"], 5.63)
        self.assertEqual(totals["Grand Total"], 4.94)

    def test_wrong_depot_mode_or_last_date_cannot_be_cached(self):
        html = (FIXTURES / "PRODDUTUR/2026-05-UD.html").read_text()
        for depot, mode, through in (("RAJAMPET", "UD", None), ("PRODDUTUR", "UM", None),
                                    ("PRODDUTUR", "UD", date(2026, 5, 30))):
            with self.subTest(depot=depot, mode=mode):
                with self.assertRaises(ValueError):
                    d.parse_engine(html, depot, "2026-05", mode, through)
        with self.assertRaises(ValueError):
            d.parse_engine("<h1>Login</h1><table><tr><td>Staff No</td></tr></table>",
                           "PRODDUTUR", "2026-05", "UD")

    def test_zero_blank_percent_and_non_numeric_are_distinct(self):
        for value, expected in (("0", 0), (0, 0), ("", None), ("—", None), ("12.50", 12.5), ("12.50%", 12.5)):
            self.assertEqual(d.number(value), expected)
        for value in ("MANUAL", "5.2 kmpl", "NaN", "-1"):
            with self.assertRaises(ValueError):
                d.number(value)

    def test_real_product_column_changes(self):
        may = engine("PRODDUTUR", "2026-05", "UM")
        july = engine("PRODDUTUR", "2026-07", "UM")
        self.assertNotIn("SUP-LUX-40-HB", may["headers"])
        self.assertIn("SUP-LUX-40-HB", july["headers"])

    def test_exact_total_size_tyre_rows_and_source_headers(self):
        for depot, code in (("PRODDUTUR", "PDT"), ("RAJAMPET", "RJP")):
            for group, columns in (("B", 24), ("F", 18)):
                html = (FIXTURES / depot / f"2026-05-{group}.html").read_text()
                snap = d.parse_tyre(html, group, code, "2026-05", depot)
                self.assertEqual(len(snap["headers"]), columns)
                values = dict(zip(snap["headers"], snap["rows"][0]))
                self.assertEqual(values["Depot"], code)
                self.assertEqual(values["Tyre size"], "All Tyre Sizes Total")
                if depot == "PRODDUTUR" and group == "B":
                    self.assertEqual(values["Num. of Tyres received"], 32)
                    self.assertEqual(values["Total Mech Defects"], 12)
                    self.assertEqual(values["Mech. Defetcs %"], 37.5)
                    self.assertEqual(values["Total premature failure"], 7)
                if depot == "PRODDUTUR" and group == "F":
                    self.assertEqual(values["Number of Tyres Received"], 73)
                    self.assertEqual(values["Repair"], 10)
                    self.assertEqual(values["Total RC"], 32)
                    self.assertEqual(values["Worn Smooth %"], 12.5)
                for wrong_code, wrong_period in (("WRONG", "2026-05"), (code, "2026-06")):
                    with self.assertRaises(ValueError):
                        d.parse_tyre(html, group, wrong_code, wrong_period, depot)

    def test_tyre_adapter_sends_total_size_and_depot_code_not_vehicle_code(self):
        session = Mock()
        session.get.return_value.text = (FIXTURES / "PRODDUTUR/2026-05-B.html").read_text()
        route = Mock(return_value=("PDT", "KADAPA(KDP ZONE)", "DPTO YSR KADAPA"))
        snap = d.SourceAdapter(session, "PRODDUTUR", "YSRKADAPA", route,
                               today=date(2026, 10, 3))("B", "2026-05")
        route.assert_called_once_with("PRODDUTUR", 2026, 5)
        session.get.assert_called_once_with(d.BASE + "/tyres/b_statement_final.php",
            params={"zone": "", "region": "", "depot": "PDT", "month_year": "MAY-2026",
                    "tyre_size": "All Tyre Sizes Total"}, timeout=45)
        self.assertFalse(snap["provisional"])

    def test_source_adapter_posts_explicit_mode_date_and_depot(self):
        session = Mock()
        session.post.return_value.text = (FIXTURES / "PRODDUTUR/2026-05-UD.html").read_text()
        adapter = d.SourceAdapter(session, "PRODDUTUR", "YSRKADAPA", Mock(), today=date(2026, 10, 3))
        snap = adapter("UD", "2026-05")
        session.post.assert_called_once_with(d.BASE + "/mednew/eng_prod_um.php",
            data={"fdate": "31/5/2026", "reg": "YSRKADAPA", "dept": "PRODDUTUR", "fupto": "UD"}, timeout=45)
        self.assertFalse(snap["provisional"])

    def test_open_month_cutoff_is_yesterday_not_future_month_end(self):
        self.assertEqual(d.effective_end("2026-10", date(2026, 10, 3)), date(2026, 10, 2))
        self.assertEqual(d.effective_end("2026-05", date(2026, 10, 3)), date(2026, 5, 31))
        with self.assertRaises(ValueError):
            d.effective_end("2026-10", date(2026, 10, 1))
        with self.assertRaises(ValueError):
            d.effective_end("2026-11", date(2026, 10, 3))


class HistoryAndViewTests(unittest.TestCase):
    def test_checkpoint_failure_stops_before_later_month_or_um(self):
        cache, fetch = d.new_cache("PRODDUTUR"), Mock(side_effect=fixture)
        checkpoint = Mock(side_effect=OSError("cache write failed"))
        with self.assertRaises(OSError):
            d.update(cache, ["2026-05", "2026-06"], fetch, selected_um="2026-06", checkpoint=checkpoint)
        self.assertEqual(fetch.call_count, 4)
        self.assertEqual(set(cache["months"]), {"2026-05"})

    def test_selected_um_replaces_rows_and_columns_that_disappear(self):
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05"], fixture, selected_um="2026-05")
        old_month = copy.deepcopy(cache["months"]["2026-05"]["UD"])
        fresh = fixture("UM", "2026-05")
        fresh["headers"] = ["SNO", "ENGINE MODEL", "Grand Total"]
        fresh["rows"] = [[1, "AL BSVI", 5.62], ["TOTAL", "", 4.98]]
        d.update(cache, ["2026-05"], lambda g, p: fresh, selected_um="2026-05")
        self.assertEqual(cache["months"]["2026-05"]["UM"], fresh)
        self.assertEqual(cache["months"]["2026-05"]["UD"], old_month)

    def test_may_july_may_preserves_closed_history_and_refreshes_only_selected_um(self):
        for depot in ("PRODDUTUR", "RAJAMPET"):
            cache = d.new_cache(depot)
            fetch = Mock(side_effect=lambda g, p: fixture(g, p, depot))
            d.update(cache, d.annual_periods("2026-05"), fetch, selected_um="2026-05")
            self.assertEqual(fetch.call_count, 14 * 4 + 1)
            closed = copy.deepcopy({p: v for p, v in cache["months"].items() if p < "2026-04"})
            april_may = copy.deepcopy({p: cache["months"][p]["UD"] for p in ("2026-04", "2026-05")})
            fetch.reset_mock()
            d.update(cache, d.annual_periods("2026-07"), fetch, selected_um="2026-07")
            self.assertEqual(fetch.call_count, 9)
            self.assertEqual({c.args[1] for c in fetch.call_args_list}, {"2026-06", "2026-07"})
            july = copy.deepcopy(cache["months"]["2026-07"])
            fetch.reset_mock()
            d.update(cache, d.annual_periods("2026-05"), fetch, selected_um="2026-05")
            fetch.assert_called_once_with("UM", "2026-05")
            self.assertEqual(cache["months"]["2026-07"], july)
            self.assertEqual({p: v for p, v in cache["months"].items() if p < "2026-04"}, closed)
            self.assertEqual({p: cache["months"][p]["UD"] for p in april_may}, april_may)
            self.assertEqual(d.sections(cache, "2026-05", True)[d.ENGINE_TITLE][0]["rows"],
                             engine(depot, "2026-05", "UM")["rows"])

    def test_failed_refresh_preserves_the_complete_previous_matrix_and_reports_staleness(self):
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05"], fixture, selected_um="2026-05")
        previous = copy.deepcopy(cache["months"])
        fetch = Mock(side_effect=TimeoutError("source unavailable"))
        d.update(cache, ["2026-05"], fetch, selected_um="2026-05")
        self.assertEqual(cache["months"], previous)
        self.assertIn("UM", cache["errors"]["2026-05"])
        self.assertEqual(d.sections(cache, "2026-05", True)[d.ENGINE_TITLE][0]["stale"], ["2026-05"])

    def test_corrupt_missing_chunks_wrong_depot_and_bad_shapes_fail_closed(self):
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05"], fixture)
        cache["metadata"] = "x" * 35000
        rows = d.encode(cache)
        self.assertEqual(d.decode(rows + [["ignored stale chunk"]], "PRODDUTUR"), cache)
        for damaged, depot in ((rows[:-1], "PRODDUTUR"), (rows, "RAJAMPET"),
                               ([rows[0], ["wrong"]], "PRODDUTUR")):
            with self.assertRaises(ValueError):
                d.decode(damaged, depot)
        bad = copy.deepcopy(cache)
        bad["months"]["2026-05"]["UD"]["rows"][0].pop()
        with self.assertRaises(ValueError):
            d.decode(d.encode(bad), "PRODDUTUR")

    def test_current_month_snapshots_refresh_then_become_closed_history(self):
        cache = d.new_cache("PRODDUTUR")
        source = Mock(side_effect=lambda g, p: dict(fixture(g, p), provisional=True))
        d.update(cache, ["2026-05"], source)
        d.update(cache, ["2026-05"], source)
        self.assertEqual(source.call_count, 8)
        source.reset_mock()
        source.side_effect = fixture
        d.update(cache, ["2026-05"], source)
        self.assertEqual(source.call_count, 4)
        source.reset_mock()
        d.update(cache, ["2026-05"], source)
        source.assert_not_called()

    def test_current_tyre_statements_do_not_claim_a_daily_cutoff(self):
        cache, wb = d.new_cache("PRODDUTUR"), Workbook()
        d.update(cache, ["2026-05"], lambda g, p: dict(fixture(g, p), provisional=True))
        d.render_tabs(wb, cache, "2026-05")
        notes = [c.value for row in wb[d.TYRE_TITLE] for c in row if c.value is not None]
        self.assertEqual(sum("no daily cutoff supplied" in str(v) for v in notes), 3)
        self.assertFalse(any("through completed days" in str(v) for v in notes))

    def test_monthly_refresh_uses_same_source_period_but_can_update_that_month(self):
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05"], fixture)
        fetch = Mock(side_effect=fixture)
        d.update(cache, ["2026-05"], fetch, refresh_monthly=True)
        self.assertEqual(fetch.call_args_list[3].args, ("UD", "2026-05"))
        self.assertEqual(fetch.call_count, 4)
        self.assertNotIn("UM", cache["months"]["2026-05"])

    def test_no_fy24_source_fetches(self):
        cache, fetch = d.new_cache("PRODDUTUR"), Mock()
        d.update(cache, ["2024-06", "2025-03"], fetch, selected_um="2025-03")
        fetch.assert_not_called()
        self.assertEqual(d.annual_periods("2025-03"), [])

    def test_monthly_selection_is_not_restricted_by_annual_history_floor(self):
        cache, fetch = d.new_cache("PRODDUTUR"), Mock(side_effect=fixture)
        d.update(cache, ["2024-06"], fetch, refresh_monthly=True)
        self.assertEqual(fetch.call_count, 4)
        self.assertEqual({c.args for c in fetch.call_args_list}, {("B", "2024-06"), ("C", "2024-06"), ("F", "2024-06"), ("UD", "2024-06")})
        blocks = d.sections(cache, "2024-06", False)
        self.assertEqual(len(blocks[d.ENGINE_TITLE]), 1)
        self.assertIn("June 2024", blocks[d.ENGINE_TITLE][0]["title"])
        self.assertFalse(blocks[d.ENGINE_TITLE][0]["missing"])

    def test_july_to_may_removes_future_rows_columns_merges_images_and_formats(self):
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05", "2026-07"], fixture, selected_um="2026-07")
        wb = Workbook()
        original = wb.active
        original.title = "Existing KPI Data"
        original["A1"], original["B1"] = 17, "=A1*2"
        d.render_tabs(wb, cache, "2026-07", annual=True)
        july = wb[d.ENGINE_TITLE]
        july_rows, july_cols = july.max_row, july.max_column
        self.assertTrue(any(c.value == "SUP-LUX-40-HB" for row in july for c in row))
        d.render_tabs(wb, cache, "2026-05", annual=True)
        may = wb[d.ENGINE_TITLE]
        self.assertLess(may.max_row, july_rows)
        self.assertLess(may.max_column, july_cols)
        self.assertFalse(any(c.value == "SUP-LUX-40-HB" for row in may for c in row))
        self.assertEqual(original["B1"].value, "=A1*2")
        self.assertEqual(len(may._images), 1)
        self.assertTrue(all(m.max_col <= may.max_column and m.max_row <= may.max_row for m in may.merged_cells.ranges))
        d.write_xlsx_cache(wb, cache)
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        saved = load_workbook(stream)
        self.assertEqual(saved[d.CACHE_TITLE].sheet_state, "hidden")
        self.assertEqual(d.decode(list(saved[d.CACHE_TITLE].values), "PRODDUTUR"), cache)

    def test_google_replacement_shrinks_owned_grid_and_clears_stale_formatting(self):
        cache, wb = d.new_cache("PRODDUTUR"), Workbook()
        d.update(cache, ["2026-05"], fixture, selected_um="2026-05")
        d.render_tabs(wb, cache, "2026-05", annual=True)
        sheet = wb[d.ENGINE_TITLE]
        before = sheet.max_row
        old = {"conditionalFormats": [{}, {}], "bandedRanges": [{"bandedRangeId": 123}],
               "charts": [{"chartId": 456}]}
        requests = d.google_requests(sheet, 77, old)
        self.assertEqual(sheet.max_row, before)
        resized = next(r["updateSheetProperties"] for r in requests
                       if "rowCount" in r.get("updateSheetProperties", {}).get("properties", {}).get("gridProperties", {}))
        self.assertEqual(resized["properties"]["gridProperties"]["rowCount"], before + 2)
        self.assertEqual(resized["properties"]["gridProperties"]["columnCount"], sheet.max_column)
        replace = next(r["updateCells"] for r in requests if "updateCells" in r)
        self.assertIn("userEnteredFormat", replace["fields"])
        self.assertIn("userEnteredValue", replace["fields"])
        self.assertEqual([r["deleteConditionalFormatRule"]["index"] for r in requests if "deleteConditionalFormatRule" in r], [1, 0])
        for request in requests:
            for kind in ("updateCells", "unmergeCells", "mergeCells"):
                if kind in request:
                    self.assertEqual(request[kind]["range"]["sheetId"], 77)

    def test_small_depot_matrix_has_no_unused_product_columns_and_preserves_colour_rules(self):
        cache, wb = d.new_cache("RAJAMPET"), Workbook()
        cache["months"]["2026-05"] = {"UD": engine("RAJAMPET", "2026-05", "UD")}
        d.render_tabs(wb, cache, "2026-05")
        sheet = wb[d.ENGINE_TITLE]
        self.assertEqual(sheet.max_column, 7)
        self.assertEqual(sheet["C10"].value, 4.69)
        self.assertEqual(sheet["C10"].fill.fgColor.rgb[-6:], "F4CCCC")
        self.assertEqual(sheet["D9"].value, 5.16)
        self.assertEqual(sheet["D9"].fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertIsNone(sheet["C9"].value)
        self.assertEqual(len(sheet._images), 1)

    def test_monthly_attach_reads_only_same_depot_month_cache_and_preserves_core(self):
        wb = Workbook()
        wb.active.title = "Monthly KMPL"
        wb.active["A1"] = "original vehicle report"
        cache = d.new_cache("PRODDUTUR")
        d.update(cache, ["2026-05"], fixture)
        api, source = Mock(), Mock(side_effect=fixture)
        with patch("src.integrations.google_drive.find_file", return_value={"id": "same-month-id"}) as find, \
                patch("src.integrations.google_sheets.sheets_service", return_value=api), \
                patch.object(d, "read_google_cache", return_value=cache) as read, \
                patch.object(d, "SourceAdapter", return_value=source):
            d.attach_monthly_tabs(wb, Mock(), "PRODDUTUR", "YSRKADAPA", "2026-05", "monthly-folder")
        find.assert_called_once_with("monthly-folder", "PRODDUTUR_2026-05")
        read.assert_called_once_with(api, "same-month-id", "PRODDUTUR")
        self.assertEqual(source.call_count, 4)
        self.assertEqual(wb["Monthly KMPL"]["A1"].value, "original vehicle report")
        self.assertTrue(set(d.TAB_TITLES) <= set(wb.sheetnames))
        self.assertEqual(wb[d.CACHE_TITLE].sheet_state, "hidden")

    def test_monthly_cache_read_failure_stops_without_replacing_views(self):
        wb = Workbook()
        wb.active["A1"] = "keep original"
        with patch("src.integrations.google_drive.find_file", return_value={"id": "same-month-id"}), \
                patch("src.integrations.google_sheets.sheets_service"), \
                patch.object(d, "read_google_cache", side_effect=ValueError("checksum")), \
                patch.object(d, "SourceAdapter") as adapter:
            with self.assertRaises(ValueError):
                d.attach_monthly_tabs(wb, Mock(), "PRODDUTUR", "YSRKADAPA", "2026-05", "monthly-folder")
        adapter.assert_not_called()
        self.assertEqual(wb.active["A1"].value, "keep original")
        self.assertFalse(set(d.TAB_TITLES) & set(wb.sheetnames))


if __name__ == "__main__":
    unittest.main()
