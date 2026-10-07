import ast
import copy
from datetime import date
from pathlib import Path
import unittest
import tempfile
from unittest.mock import Mock, patch

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
import kmpl_ranges as k
import report_details as d

FIXTURES = Path(__file__).parent / "fixtures"


def html(entity="VEHICLE",period="2026-09"):
    raw=(FIXTURES/f"kmpl-range-{entity.lower()}-2026-09.html").read_text()
    return raw.replace("09/2026",d.month_end(period).strftime("%m/%Y")).replace("202609",period.replace("-",""))


def snap(entity="VEHICLE",period="2026-09",depot="PRODDUTUR"):
    return k.parse_range(html(entity,period),entity,depot,period,"YSRKADAPA")


class RangeTests(unittest.TestCase):
    def test_owner_source_counts_and_regional_controls(self):
        for entity,counts,total in [("VEHICLE",[0,2,39,18,9,9,4,81],411),
                                    ("DRIVER",[0,3,78,36,35,16,5,173],911)]:
            s=snap(entity)
            self.assertEqual(s["counts"],counts)
            self.assertEqual(s["regional_total"][-1],total)
            self.assertEqual(snap(entity,depot="RAJAMPET")["counts"][-1],43 if entity=="VEHICLE" else 82)

    def test_reject_wrong_month_entity_login_missing_and_duplicate_table(self):
        for text in [html().replace("09/2026","08/2026"),html("DRIVER"),"<form>Login</form>",html()+html(),html().replace("PRODDUTUR","OTHER")]:
            with self.subTest(text=text[:60]), self.assertRaises(ValueError):
                k.parse_range(text,"VEHICLE","PRODDUTUR","2026-09","YSRKADAPA")

    def test_reject_ambiguous_labels_blank_fraction_negative_counts(self):
        for text in [html().replace("5.60 - ABV","5.61 - ABV"),html().replace("<td>0</td>","<td></td>",1),
                     html().replace("<td>0</td>","<td>-1</td>",1),html().replace("<td>0</td>","<td>0.5</td>",1)]:
            with self.subTest(text=text[:60]),self.assertRaises(ValueError):
                k.parse_range(text,"VEHICLE","PRODDUTUR","2026-09","YSRKADAPA")

    def test_reject_row_total_regional_total_and_drilldown_mismatch(self):
        for text in [html().replace("<th>81</th>","<th>80</th>"),html().replace("<th>411</th>","<th>410</th>"),
                     html().replace("rreg=PRODDUTUR","rreg=KADAPA"),html().replace("yymm=202609","yymm=202608",1)]:
            with self.subTest(text=text[:60]),self.assertRaises(ValueError):
                k.parse_range(text,"VEHICLE","PRODDUTUR","2026-09","YSRKADAPA")

    def test_duplicate_depot_and_duplicated_total_rejected(self):
        soup=BeautifulSoup(html(),"html.parser")
        soup.table.insert(3,copy.deepcopy(soup.table.find_all("tr")[2]))
        with self.assertRaises(ValueError):k.parse_range(str(soup),"VEHICLE","PRODDUTUR","2026-09","YSRKADAPA")
        soup=BeautifulSoup(html(),"html.parser");soup.table.append(copy.deepcopy(soup.table.find_all("tr")[-1]))
        with self.assertRaises(ValueError):k.parse_range(str(soup),"VEHICLE","PRODDUTUR","2026-09","YSRKADAPA")

    def test_get_endpoint_and_payload_current_is_provisional_future_rejected(self):
        for entity,path in k.ENDPOINTS.items():
            session=Mock();session.get.return_value.text=html(entity)
            result=k.SourceAdapter(session,"PRODDUTUR","YSRKADAPA",date(2026,9,20))(entity,"2026-09")
            session.get.assert_called_once_with(f"{d.BASE}/med/{path}",params={"action":"","yymm":"202609","rreg":"YSRKADAPA"},timeout=45)
            session.post.assert_not_called();self.assertTrue(result["provisional"])
            session.get.reset_mock()
            with self.assertRaises(ValueError):k.SourceAdapter(session,"PRODDUTUR","YSRKADAPA",date(2026,8,20))(entity,"2026-09")
            session.get.assert_not_called()

    def test_fy_rollover_cutoff_and_no_future_months(self):
        self.assertEqual(k.periods_for_fys(["2025-26"],"2026-03"),[f"2025-{m:02d}" for m in range(4,13)]+["2026-01","2026-02","2026-03"])
        self.assertEqual(k.periods_for_fys(["2026-27"],"2026-05"),["2026-04","2026-05"])
        with self.assertRaises(ValueError):k.periods_for_fys(["2026-27"],"2026-13")

    def test_may_july_may_reuses_closed_history_and_refreshes_selection(self):
        for depot in ["PRODDUTUR","RAJAMPET"]:
            cache=k.new_cache(depot);fetch=lambda e,p:snap(e,p,depot)
            self.assertEqual(len(k.update(cache,["2026-04","2026-05"],fetch,"2026-05")),4)
            old=copy.deepcopy(cache["months"]);wb=Workbook()
            k.render_tab(wb,cache,"2026-05",["2026-27"])
            initial=[[c.value for c in row] for row in wb[k.ANNUAL_TITLE]]
            self.assertEqual(len(k.update(cache,["2026-04","2026-05","2026-06","2026-07"],fetch,"2026-07")),4)
            self.assertEqual(cache["months"]["2026-04"],old["2026-04"])
            k.render_tab(wb,cache,"2026-07",["2026-27"]);july_rows=wb[k.ANNUAL_TITLE].max_row
            self.assertEqual(len(k.update(cache,["2026-04","2026-05"],fetch,"2026-05")),2)
            k.render_tab(wb,cache,"2026-05",["2026-27"])
            self.assertEqual(initial,[[c.value for c in row] for row in wb[k.ANNUAL_TITLE]])
            self.assertLess(wb[k.ANNUAL_TITLE].max_row,july_rows)
            self.assertIn("2026-07",cache["months"])
            self.assertNotIn("Jul-2026",[c.value for row in wb[k.ANNUAL_TITLE] for c in row])

    def test_failure_retains_same_period_snapshot_missing_is_not_zero(self):
        cache=k.new_cache("PRODDUTUR");k.update(cache,["2026-09"],snap,"2026-09")
        before=copy.deepcopy(cache["months"])
        def fail(e,p):raise TimeoutError("offline source")
        k.update(cache,["2026-09","2026-10"],fail,"2026-09")
        self.assertEqual(cache["months"]["2026-09"],before["2026-09"])
        wb=Workbook();k.render_tab(wb,cache,"2026-10",["2026-27"])
        values=[c.value for row in wb[k.ANNUAL_TITLE] for c in row]
        self.assertIn("Stale — saved source",values);self.assertIn("Unavailable — no verified source",values)
        for row in wb[k.ANNUAL_TITLE]:
            if row[0].value=="Oct-2026":self.assertTrue(all(c.value is None for c in row[1:12]))

    def test_checksum_identity_and_semantic_cache_validation(self):
        cache=k.new_cache("PRODDUTUR");k.update(cache,["2026-09"],snap,"2026-09")
        self.assertEqual(k.decode(k.encode(cache),"PRODDUTUR"),cache)
        with self.assertRaises(ValueError):k.decode(k.encode(cache),"RAJAMPET")
        rows=k.encode(cache);rows[1][0]+="x"
        with self.assertRaises(ValueError):k.decode(rows,"PRODDUTUR")
        cache["months"]["2026-09"]["VEHICLE"]["counts"][0]=False
        with self.assertRaises(ValueError):k.decode(k.encode(cache),"PRODDUTUR")

    def test_checkpoint_failure_is_not_swallowed(self):
        def fail():raise OSError("write failed")
        with self.assertRaises(OSError):k.update(k.new_cache("PRODDUTUR"),["2026-09"],snap,"2026-09",fail)

    def test_provisional_history_refetched_and_missing_entity_repaired(self):
        cache=k.new_cache("PRODDUTUR");k.update(cache,["2026-05","2026-09"],snap,"2026-09")
        cache["months"]["2026-05"]["VEHICLE"]["provisional"]=True
        del cache["months"]["2026-05"]["DRIVER"]
        self.assertEqual(len(k.update(cache,["2026-05","2026-09"],snap,"2026-09")),4)

    def test_workbook_preserves_core_formulas_cache_and_native_types(self):
        cache=k.new_cache("PRODDUTUR");k.update(cache,["2026-09"],snap,"2026-09")
        wb=Workbook();wb.active.title="Core";wb.active["A1"]="=1+1"
        ws=k.render_tab(wb,cache,"2026-09");k.write_xlsx_cache(wb,cache)
        self.assertEqual(wb["Core"]["A1"].value,"=1+1");self.assertEqual(wb[k.CACHE_TITLE].sheet_state,"hidden")
        self.assertEqual(ws["B8"].value,0);self.assertEqual(ws["I8"].value,81)
        self.assertEqual(ws["K8"].value,'=IF(I8=0,"",J8/I8)')
        self.assertEqual(ws["K8"].number_format,"0.0%")
        requests=d.google_requests(ws,17)
        records=next(r["updateCells"]["rows"] for r in requests if "updateCells" in r)
        self.assertEqual(records[7]["values"][1]["userEnteredValue"],{"numberValue":0})
        self.assertEqual(records[7]["values"][10]["userEnteredValue"],{"formulaValue":ws["K8"].value})
        self.assertEqual(ws.max_column,13);self.assertEqual(len(ws._images),1)
        self.assertIn("A3:C3",[str(m) for m in ws.merged_cells.ranges])

    def test_zero_population_retains_zero_and_empty_share_formula(self):
        cache=k.new_cache("PRODDUTUR");s=snap()
        s["counts"]=[0]*8
        for row in s["rows"]:
            if row[1]=="PRODDUTUR":row[2:]=[0]*8
        s["regional_total"]=[sum(row[i+2] for row in s["rows"]) for i in range(8)]
        k.validate_snapshot(s,"PRODDUTUR","VEHICLE","2026-09")
        cache["months"]["2026-09"]={"VEHICLE":s}
        wb=Workbook();ws=k.render_tab(wb,cache,"2026-09")
        self.assertEqual(ws["I8"].value,0);self.assertIn('IF(I8=0,""',ws["K8"].value)

    def test_both_entry_points_and_annual_native_visibility(self):
        monthly=Path("monthly_vehicle_report.py").read_text();annual=Path("annual_kpi_runner_v11.py").read_text()
        self.assertIn("kmpl_ranges.attach_monthly(workbook, session, display_name, region_code, args.month, folder_id)",monthly)
        tree=ast.parse(annual);finalizer=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=="_finalize_live_workbook")
        self.assertIn("ranges.ANNUAL_TITLE",ast.unparse(finalizer))
        self.assertIn("ranges.publish_google_tab",annual)

    def test_saved_fy_charts_have_meaningful_labels_native_ranges_and_gaps(self):
        cache=k.new_cache("PRODDUTUR");k.update(cache,["2026-09"],snap,"2026-09")
        wb=Workbook();k.render_tab(wb,cache,"2026-09",["2024-25","2025-26","2026-27"])
        with tempfile.TemporaryDirectory() as path:
            file=Path(path)/"charts.xlsx";wb.save(file)
            ws=load_workbook(file)[k.ANNUAL_TITLE]
            self.assertEqual(len(ws._charts),2)
            self.assertEqual([s.tx.v for s in ws._charts[0].series],["First 3 ranges %","Highest range %"])
            requests=k.google_chart_requests(ws,17)
            self.assertEqual(len(requests),2)
            for request in requests:
                basic=request["addChart"]["chart"]["spec"]["basicChart"]
                self.assertFalse(basic["interpolateNulls"])
                domain=basic["domains"][0]["domain"]["sourceRange"]["sources"][0]
                values=basic["series"][0]["series"]["sourceRange"]["sources"][0]
                self.assertEqual(values["startRowIndex"],domain["startRowIndex"])
                self.assertEqual(values["endRowIndex"],domain["endRowIndex"])
                self.assertEqual(ws.cell(domain["startRowIndex"]+1,1).value,"Month")



def test_month_and_upto_cache_are_independent():
    cache = k.new_cache("Proddutur")
    calls = []
    def fetch(entity, period, scope="MONTH"):
        calls.append((entity, period, scope))
        counts = [1,2,3,4,5,6,7,28] if scope == "MONTH" else [2,3,4,5,6,7,8,35]
        return dict(depot="Proddutur", period=period, entity=entity, scope=scope, region="YSRKADAPA",
                    headers=["SL No","Depot",*k.LABELS,"G.Total"],
                    rows=[[1,"Proddutur",*counts]], regional_total=counts, counts=counts,
                    source="fixture", provisional=False)
    k.update(cache, ["2026-09"], fetch, "2026-09", include_upto=True)
    groups = cache["months"]["2026-09"]
    assert groups["VEHICLE_MONTH"]["counts"][-1] == 28
    assert groups["VEHICLE_UPTO"]["counts"][-1] == 35
    assert groups["DRIVER_MONTH"]["counts"][-1] == 28
    assert groups["DRIVER_UPTO"]["counts"][-1] == 35
    assert len(calls) == 4


def test_monthly_combined_table_and_annual_upto_row():
    from openpyxl import Workbook
    cache = k.new_cache("Proddutur")
    month=[1,2,3,4,5,6,7,28]
    upto=[2,3,4,5,6,7,8,35]
    for entity in k.ENDPOINTS:
        cache["months"].setdefault("2026-09",{})[f"{entity}_MONTH"] = dict(
            depot="Proddutur",period="2026-09",entity=entity,scope="MONTH",region="YSRKADAPA",
            headers=["SL No","Depot",*k.LABELS,"G.Total"],rows=[[1,"Proddutur",*month]],
            regional_total=month,counts=month,source="fixture",provisional=False)
        cache["months"]["2026-09"][f"{entity}_UPTO"] = dict(
            depot="Proddutur",period="2026-09",entity=entity,scope="UPTO",region="YSRKADAPA",
            headers=["SL No","Depot",*k.LABELS,"G.Total"],rows=[[1,"Proddutur",*upto]],
            regional_total=upto,counts=upto,source="fixture",provisional=False)
    wb=Workbook(); del wb[wb.sheetnames[0]]
    ws=k.render_tab(wb,cache,"2026-09")
    values=[[c.value for c in row] for row in ws.iter_rows()]
    assert ["KMPL Range","Vehicle Month","Vehicle Upto","Driver Month","Driver Upto"] in [r[:5] for r in values]
    wb2=Workbook(); del wb2[wb2.sheetnames[0]]
    ws2=k.render_tab(wb2,cache,"2026-09",["2026-27"])
    assert any(c.value=="Upto/Cum Sep-26" for row in ws2.iter_rows() for c in row)


def test_upto_source_contracts_use_verified_selector_and_result_endpoints():
    for entity in k.ENDPOINTS:
        session=Mock()
        response=Mock(); response.text=html(entity)
        session.get.return_value=response
        result=k.SourceAdapter(session,"PRODDUTUR","YSRKADAPA",date(2026,10,7))(entity,"2026-09","UPTO")
        assert result["scope"]=="UPTO"
        assert session.get.call_args_list[0].args[0] == f"{d.BASE}/med/{k.UPTO_SELECTORS[entity]}"
        assert session.get.call_args_list[1].args[0] == f"{d.BASE}/med/{k.UPTO_ENDPOINTS[entity]}"
        assert session.get.call_args_list[1].kwargs["params"] == {"action":"","yymm":"202609","rreg":"YSRKADAPA"}

if __name__=="__main__":
    unittest.main()
