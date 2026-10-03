"""Read-only source review; generate isolated range-tab XLSX previews.

No Drive writes, Google Sheets updates, Hub callbacks or Telegram delivery.
--offline uses only the owner-supplied September tables, without invented history.
"""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import kmpl_ranges as k


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--offline",action="store_true")
    parser.add_argument("--probe-selection",action="store_true",help="Read-only GET/POST request contract probe")
    parser.add_argument("--month",default="2026-09")
    parser.add_argument("--output",default="reports/range-review")
    args=parser.parse_args()
    k.d.month_end(args.month)
    if args.offline and args.month!="2026-09":
        parser.error("Owner source fixtures exist only for September 2026")
    output=Path(args.output);output.mkdir(parents=True,exist_ok=True)
    if not args.offline:
        from src.auth.client import login
        session=login()
        if args.probe_selection:
            from bs4 import BeautifulSoup
            payload=dict(action="",yymm=args.month.replace("-",""),rreg="YSRKADAPA")
            for entity,endpoint in k.ENDPOINTS.items():
                url=f"{k.d.BASE}/med/{endpoint}"
                for method in ["GET","POST","POST_QUERY"]:
                    response=(session.get(url,params=payload,timeout=45) if method=="GET" else
                              session.post(url,data=payload,params=payload if method=="POST_QUERY" else None,timeout=45))
                    response.raise_for_status()
                    soup=BeautifulSoup(response.text,"html.parser")
                    print("RANGE_SELECTION_PROBE",entity,method,"table_rows",[len(t.find_all('tr')) for t in soup.find_all('table')])
                    try:
                        result=k.parse_range(response.text,entity,"PRODDUTUR",args.month,"YSRKADAPA")
                        print("RANGE_SELECTION_VALID",entity,method,result["counts"])
                    except Exception as exc:
                        print("RANGE_SELECTION_INVALID",entity,method,str(exc)[:140])
            return
    summary={"mode":"owner source fixtures" if args.offline else "authenticated read-only source review","depots":{}}
    first_year=k.d.fy_start(args.month).year
    fys=[f"{year}-{str(year+1)[-2:]}" for year in range(first_year-2,first_year+1)]
    for depot in ["PRODDUTUR","RAJAMPET"]:
        cache=k.new_cache(depot)
        if args.offline:
            def fetch(entity,period):
                text=Path(f"tests/fixtures/kmpl-range-{entity.lower()}-2026-09.html").read_text()
                return k.parse_range(text,entity,depot,period,"YSRKADAPA")
            periods=[args.month]
        else:
            fetch=k.SourceAdapter(session,depot,"YSRKADAPA")
            periods=sorted(set(p for p in [f"{first_year-2}-04",f"{first_year-1}-04",f"{first_year}-05",f"{first_year}-07",args.month] if p<=args.month))
        k.update(cache,periods,fetch,args.month)
        if any(entity not in cache["months"].get(args.month,{}) for entity in k.ENDPOINTS):
            raise RuntimeError(f"{depot} selected month has unavailable range sources: {cache['errors']}")
        for annual in [False,True]:
            wb=Workbook();wb.remove(wb.active)
            k.render_tab(wb,cache,args.month,fys if annual else None)
            k.write_xlsx_cache(wb,cache)
            path=output/f"{depot}_{args.month}_{'annual' if annual else 'monthly'}_ranges.xlsx"
            wb.save(path)
            saved=load_workbook(path)
            assert k.decode(list(saved[k.CACHE_TITLE].values),depot)==cache
            if annual:
                requests=k.google_chart_requests(saved[k.ANNUAL_TITLE],17)
                assert len(requests)==len(saved[k.ANNUAL_TITLE]._charts)
                assert len(requests)>=2
                assert all(r['addChart']['chart']['spec']['basicChart']['headerCount']==1 for r in requests)
            else:
                assert not saved[k.MONTHLY_TITLE]._charts
        summary["depots"][depot]={"selected":{e:cache["months"][args.month][e]["counts"] for e in k.ENDPOINTS},
                                  "sampled_periods":periods,"source_errors":cache["errors"]}
        (output/f"{depot}_range_cache.json").write_text(json.dumps(cache,indent=2))
    (output/"summary.json").write_text(json.dumps(summary,indent=2))
    print(json.dumps(summary,indent=2))


if __name__=="__main__":main()
