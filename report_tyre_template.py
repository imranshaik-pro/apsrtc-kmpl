"""Owner's template. Raw B/C/F snapshots stay immutable; presentation is derived.

B mechanical defects plus F's two percentages, B premature failures, then C's
four grouped stages. Receipts and F numerator/denominator counts are hidden.
"""
from openpyxl.utils import get_column_letter
import report_details as d

DEFECTS = ("Misalignment", "Mismatching", "Camber wear", "King pin play", "Brake Binding",
           "Over Inflation", "Under Inflation", "Hub Bearing play")
FAILURES = ("S1", "S2", "S3", "S6", "S8", "S9")
INPUT_HEADERS = ("Month", "B Tyres Received", "F Tyres Received", "F Stones",
                 "F Worn Smooth RC", "F Worn Smooth RT", "F Total RC", "F Total RT")


def metrics(cache, period, group):
    snap = cache["months"].get(period, {}).get(group)
    return {d.key(h): v for h, v in zip(snap["headers"], snap["rows"][0])} if snap else {}


def value(values, label):
    return values.get(d.key(label))


def ratio(values, label):
    x = value(values, label)
    return x / 100 if isinstance(x, (int, float)) else None


def write_inputs(workbook, cache, periods):
    if d.TYRE_INPUT_TITLE in workbook.sheetnames:
        del workbook[d.TYRE_INPUT_TITLE]
    ws = workbook.create_sheet(d.TYRE_INPUT_TITLE)
    ws.append(INPUT_HEADERS)
    for period in periods:
        b, f = metrics(cache, period, "B"), metrics(cache, period, "F")
        ws.append([period, value(b, "Num. of Tyres received"), value(f, "Number of Tyres Received"),
                   value(f, "Stones"), value(f, "Worn smooth RC"), value(f, "Worn smooth RT"),
                   value(f, "Total RC"), value(f, "Total RT")])
    ws.sheet_state = "hidden"


def scrap_notes(cache, periods):
    notes = []
    for period in periods:
        c = metrics(cache, period, "C")
        for stage in d.SCRAP_GROUPS:
            codes = [value(c, f"{stage} — S{i}") for i in range(1, 10)]
            reported = value(c, f"{stage} — Total")
            if isinstance(reported, (int, float)) and all(isinstance(x, (int, float)) for x in codes):
                if sum(codes) != reported:
                    notes.append(f"{d.month_end(period).strftime('%b %Y')} · {stage}: source total {reported}; "
                                 f"S1–S9 sum {sum(codes)}. Source total retained.")
    return notes


def tyre_sections(cache, periods, annual):
    input_rows = {p: i + 2 for i, p in enumerate(periods)}
    batches = {}
    for period in periods:
        batches.setdefault(d.fy_label(period) if annual else "", []).append(period)
    result = []
    for kind in ("mechanical", "premature", "scrap"):
        for fy, months in batches.items():
            suffix = f" · FY {fy}" if annual else ""
            if kind == "mechanical":
                groups = ("B", "F")
                headers = ["S.No", "Month", *DEFECTS, "Total Mech Defects", "Mech. Defects %", "Stone %", "Worn Smooth %"]
                title = "Mechanical defects and tyre indicators — Statements B / F" + suffix
            elif kind == "premature":
                groups = ("B",)
                headers = ["S.No", "Month", *FAILURES, "Total premature failure"]
                title = "Premature tyre failures — Statement B" + suffix
            else:
                groups = ("C",)
                headers = ["S.No", "Month"] + ([f"S{i}" for i in range(1, 10)] + ["Total"]) * 4
                title = "Scrap tyres — Statement C" + suffix
            rows = []
            for serial, period in enumerate(months, 1):
                b, f, c = [metrics(cache, period, g) for g in ("B", "F", "C")]
                row = [serial, d.month_end(period).strftime("%b-%y")]
                if kind == "mechanical":
                    row += [value(b, label) for label in (*DEFECTS, "Total Mech Defects")]
                    row += [ratio(b, "Mech. Defetcs %"), ratio(f, "Stone %"), ratio(f, "Worn Smooth %")]
                elif kind == "premature":
                    row += [value(b, label) for label in (*FAILURES, "Total premature failure")]
                else:
                    row += [value(c, f"{stage} — {label}") for stage in d.SCRAP_GROUPS
                            for label in [f"S{i}" for i in range(1, 10)] + ["Total"]]
                rows.append(row)
            block = dict(kind="tyre", template_kind=kind, title=title, headers=headers, rows=rows, periods=months,
                         missing=[f"{g}:{p}" for p in months for g in groups if not cache["months"].get(p, {}).get(g)],
                         stale=[f"{g}:{p}" for p in months for g in groups if cache["errors"].get(p, {}).get(g)
                                and cache["months"].get(p, {}).get(g)],
                         provisional=any(cache["months"].get(p, {}).get(g, {}).get("provisional") for p in months for g in groups),
                         source=" · ".join(f"{d.BASE}/tyres/{g.lower()}_statement_final.php" for g in groups),
                         period_text=f"{d.month_end(months[0]).strftime('%B %Y')} to {d.month_end(months[-1]).strftime('%B %Y')}",
                         percent_columns=[12, 13, 14] if kind == "mechanical" else [],
                         notes=scrap_notes(cache, months) if kind == "scrap" else [])
            if kind == "scrap":
                block["header_groups"] = [(3+i*10, 12+i*10, stage) for i, stage in enumerate(d.SCRAP_GROUPS)]
            elif kind == "mechanical":
                block["header_groups"] = [(i, i, label) for i, label in enumerate(
                    ("MA", "MM", "CBW", "KPP", "BB", "OI", "UI", "HBP"), 3)]
            if annual:
                block.update(total_index=len(rows), input_start=input_rows[months[0]], input_end=input_rows[months[-1]])
                rows.append([None, "FY Total"] + [None] * (len(headers)-2))
            result.append(block)
    return result


def finish_totals(ws, block, data_start):
    """Visible totals are editable formulas; no missing source becomes a zero."""
    if "total_index" not in block:
        return
    count = block["total_index"]
    last = data_start + count - 1
    total = last + 1
    for col in range(3, len(block["headers"]) + 1):
        letter = get_column_letter(col)
        target = f"{letter}{data_start}:{letter}{last}"
        ws.cell(total, col, f'=IF(COUNT({target})={count},SUM({target}),"")')
        ws.cell(total, col).number_format = "#,##0"
    if block["template_kind"] != "mechanical":
        return
    start, end = block["input_start"], block["input_end"]
    inputs = lambda col: f"'{d.TYRE_INPUT_TITLE}'!{col}{start}:{col}{end}"
    defects, receipts = f"K{data_start}:K{last}", inputs("B")
    ws.cell(total, 12, f'=IF(AND(COUNT({defects},{receipts})={2*count},SUM({receipts})>0),SUM({defects})/SUM({receipts}),"")')
    stones, received = inputs("D"), inputs("C")
    ws.cell(total, 13, f'=IF(AND(COUNT({stones},{received})={2*count},SUM({received})>0),SUM({stones})/SUM({received}),"")')
    worn_rc, worn_rt, total_rc, total_rt = [inputs(col) for col in "EFGH"]
    ws.cell(total, 14, f'=IF(AND(COUNT({worn_rc},{worn_rt},{total_rc},{total_rt})={4*count},SUM({total_rc},{total_rt})>0),SUM({worn_rc},{worn_rt})/SUM({total_rc},{total_rt}),"")')
    for col in (12, 13, 14):
        ws.cell(total, col).number_format = "0.00%"
